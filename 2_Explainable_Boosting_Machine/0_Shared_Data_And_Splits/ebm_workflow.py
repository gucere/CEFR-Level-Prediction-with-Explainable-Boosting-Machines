import argparse
import csv
import hashlib
import importlib.metadata
import json
import os
import re
import sys
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from contextlib import contextmanager
from pathlib import Path
import duckdb
import joblib
import numpy as np
import pandas as pd
import psutil
from interpret.glassbox import ExplainableBoostingClassifier, ExplainableBoostingRegressor
from sklearn.tree import DecisionTreeClassifier
from sklearn.dummy import DummyClassifier, DummyRegressor
from scipy import sparse
from scipy.stats import spearmanr
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, f1_score, cohen_kappa_score, mean_absolute_error, mean_squared_error, r2_score, mutual_info_score
from sklearn.utils.class_weight import compute_sample_weight
HERE = Path(__file__).resolve().parent
BASE = HERE.parent
ROOT = BASE.parent
DATA = ROOT / "1_Complexity_Analysis_Full_Module/2_Basic_Analysis/0_Feature_Dataframe"
sys.path.insert(0, str(DATA))
from create_dataframe import identifier, literal, memory_budget
SHARED = HERE
RANKING = BASE / "1_Training_Feature_Ranking"
LABELS = ["A1", "A2", "B1", "B2", "C1", "C2"]
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
DERIVED = {
    f"derived__ERRANT__{operation}_error_share": {
        "numerator": f"accuracy__ERRANT__{operation}_errors",
        "denominator": "accuracy__ERRANT__total_errors",
        "definition": f"Fraction of all correction edits classified as {operation}; zero for error-free texts.",
        "construct": "Distribution of correction operations, conditional on the total error burden",
    }
    for operation in ("missing", "unnecessary", "replacement", "other")
}

def log(message):
    print(message, flush=True)

def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))

def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    os.replace(temporary, path)

def digest(path):
    hasher = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(4 * 1024**2), b""):
            hasher.update(chunk)
    return hasher.hexdigest()

def manifest_path(path):
    path = Path(path).resolve()
    return 'Thesis/' + path.relative_to(ROOT).as_posix() if path.is_relative_to(ROOT) else str(path)


def resolve_manifest_path(value):
    value = str(value).replace('\\', '/')
    old_prefix = 'Thesis/2_Preparation_For_Experiments/3_Explainable_Boosting_Machine/'
    if value.startswith(old_prefix):
        value = 'Thesis/2_Explainable_Boosting_Machine/' + value[len(old_prefix):]
    value = value.replace('/5_LLM_Post_Training/', '/5_Data_For_LLM_Post_Training/')
    if value.startswith('Thesis/'):
        path = (ROOT / value[7:]).resolve()
        if not path.is_relative_to(ROOT):
            raise ValueError('Manifest path escapes Thesis.')
        return path
    return Path(value).resolve()


_SOURCE_HASHES = {}


def fingerprint(path):
    path = Path(path).resolve()
    stat = path.stat()
    key = (str(path), stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns)
    if key not in _SOURCE_HASHES:
        _SOURCE_HASHES[key] = digest(path)
    return {'path': manifest_path(path), 'bytes': stat.st_size, 'sha256': _SOURCE_HASHES[key]}

def token(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True).encode()).hexdigest()

def normalized_id(value):
    value = str(value).strip()
    if not re.fullmatch(r"[0-9]+(?:\.0+)?", value):
        raise ValueError("Text and learner IDs must be nonmissing integer identifiers.")
    return str(int(value.split(".")[0]))

def shared_strings(archive, wanted):
    """Decode only needed shared strings, discarding writing text as XML streams past."""
    if not wanted:
        return {}
    found, index = {}, 0
    with archive.open("xl/sharedStrings.xml") as handle:
        events = ET.iterparse(handle, events=("start", "end"))
        _, root = next(events)
        for event, element in events:
            if event == "end" and element.tag == NS + "si":
                if index in wanted:
                    found[index] = "".join(node.text or "" for node in element.iter(NS + "t"))
                index += 1
                root.clear()
                if len(found) == len(wanted):
                    break
    if set(found) != wanted:
        raise ValueError("Unresolved shared-string reference in the metadata workbook.")
    return found

def read_learner_metadata(path):
    """Read only IDs and CEFR; the source XLSX incorrectly declares dimension A1."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        frame = pd.read_csv(path, dtype=str, keep_default_na=False)
        frame = frame.rename(columns={"writing_id": "text_id", "cefr": "cefr_level"})
        frame = frame[["text_id", "learner_id", "cefr_level"]].copy()
    elif path.suffix.lower() == ".xlsx":
        records = []
        with zipfile.ZipFile(path) as archive:
            sheet = "xl/worksheets/sheet1.xml"
            with archive.open(sheet) as handle:
                for _, element in ET.iterparse(handle, events=("end",)):
                    if element.tag == NS + "row":
                        header_cells = list(element)
                        break
                else:
                    raise ValueError("Empty metadata worksheet.")
            refs = {int(c.findtext(NS + "v")) for c in header_cells if c.get("t") == "s"}
            strings = shared_strings(archive, refs)
            headers = {}
            for cell in header_cells:
                text = (strings[int(cell.findtext(NS + "v"))] if cell.get("t") == "s"
                        else "".join(t.text or "" for t in cell.iter(NS + "t")))
                headers[re.sub(r"[0-9]", "", cell.get("r"))] = text.strip().lower()
            required = {"writing_id", "learner_id", "cefr"}
            if not required.issubset(headers.values()):
                raise ValueError(f"Workbook needs columns {sorted(required)}.")
            positions = {key: name for key, name in headers.items() if name in required}
            needed_strings = set()
            with archive.open(sheet) as handle:
                events = ET.iterparse(handle, events=("start", "end"))
                _, root = next(events)
                for event, element in events:
                    if event != "end" or element.tag != NS + "row":
                        continue
                    if element.get("r") != "1":
                        row = {}
                        for cell in element:
                            name = positions.get(re.sub(r"[0-9]", "", cell.get("r", "")))
                            if name is None:
                                continue
                            kind = cell.get("t")
                            value = cell.findtext(NS + "v", "")
                            if kind == "s":
                                value = int(value)
                                needed_strings.add(value)
                            elif kind == "inlineStr":
                                value = "".join(t.text or "" for t in cell.iter(NS + "t"))
                            row[name] = (kind, value)
                        if row:
                            if set(row) != required:
                                raise ValueError("A workbook row has incomplete learner metadata.")
                            records.append(tuple(row[name] for name in ("writing_id", "learner_id", "cefr")))
                            if len(records) % 100000 == 0:
                                log(f"  Read metadata for {len(records):,} texts")
                    root.clear()
            decoded = shared_strings(archive, needed_strings)
            frame = pd.DataFrame(
                ([decoded[value] if kind == "s" else value for kind, value in row] for row in records),
                columns=["text_id", "learner_id", "cefr_level"],
            )
    else:
        raise ValueError("Metadata must be an XLSX workbook or CSV with IDs and CEFR.")
    frame["text_id"] = frame.text_id.map(normalized_id).astype(np.int64)
    frame["learner_id"] = frame.learner_id.map(normalized_id)
    frame["cefr_level"] = frame.cefr_level.str.strip().str.upper()
    if frame.text_id.duplicated().any() or not frame.cefr_level.isin(LABELS).all():
        raise ValueError("Metadata contains duplicate IDs or invalid CEFR labels.")
    return frame

def dictionary(data):
    with (data / "feature_dictionary.csv").open(encoding="utf-8-sig", newline="") as handle:
        specs = list(csv.DictReader(handle))
    names = [s["column"] for s in specs]
    if not names or len(set(names)) != len(names):
        raise ValueError("The feature dictionary must contain unique feature columns.")
    for spec in specs:
        if spec["group"] not in {"complexity", "accuracy"} or spec["column"] != f"{spec['group']}__{spec['source']}__{spec['original_feature']}":
            raise ValueError("Invalid feature dictionary mapping.")
    return specs

def assign_split(learner, seed):
    fraction = int(hashlib.sha256(f"{seed}:{learner}".encode()).hexdigest()[:16], 16) / 2**64
    return "train" if fraction < .70 else "validation" if fraction < .85 else "test"

def validate_splits(frame, labels):
    if frame.text_id.duplicated().any() or frame.learner_id.isna().any() or (frame.learner_id == "").any():
        raise ValueError("Duplicate texts or missing learner IDs in split assignments.")
    if frame.groupby("learner_id").split.nunique().max() != 1:
        raise ValueError("A learner appears in more than one split.")
    if set(frame.split) != {"train", "validation", "test"}:
        raise ValueError("Expected train, validation and test splits.")
    counts = frame.groupby(["split", "cefr_level"]).size()
    for split in ("train", "validation", "test"):
        for label in labels:
            if counts.get((split, label), 0) < 2:
                raise ValueError(f"Insufficient {label} texts in {split}. Use a larger sample or another seed.")

def feature_sql(name, originals):
    if name in originals:
        return "d." + identifier(name)
    if name not in DERIVED:
        raise ValueError(f"Unknown feature {name}; only declared measurements are allowed.")
    definition = DERIVED[name]
    numerator, denominator = ["d." + identifier(definition[key]) for key in ("numerator", "denominator")]
    return (f"CASE WHEN {denominator}=0 AND {numerator}=0 THEN 0.0 "
            f"WHEN {denominator}>0 AND {numerator} BETWEEN 0 AND {denominator} THEN {numerator}/{denominator} ELSE NULL END")

def load_matrix(manifest, rows, columns, working, name, memory):
    """Read bounded column batches into a disk-backed float64 matrix in text-ID order."""
    rows = rows.sort_values("text_id")
    result = np.lib.format.open_memmap(working / f"{name}.npy", mode="w+", dtype=np.float64, shape=(len(rows), len(columns)))
    data = resolve_manifest_path(manifest["data_directory"])
    originals = {s["column"] for s in dictionary(data)}
    with duckdb.connect() as con:
        con.execute(f"SET memory_limit={literal(memory_budget(memory))}")
        con.execute("SET threads=1")
        con.register("wanted", rows[["text_id"]])
        for start in range(0, len(columns), 16):
            batch = columns[start:start + 16]
            projection = ", ".join(f"{feature_sql(column, originals)} AS {identifier(column)}" for column in batch)
            values = con.execute(f"SELECT d.text_id, {projection} FROM read_parquet({literal(data / 'feature_dataframe.parquet')}) d JOIN wanted USING(text_id) ORDER BY d.text_id").fetchnumpy()
            if not np.array_equal(values.pop("text_id"), rows.text_id.to_numpy()):
                raise ValueError("Feature/text-ID alignment changed while loading.")
            for index, column in enumerate(batch, start):
                value = values[column]
                result[:, index] = value.filled(np.nan) if np.ma.isMaskedArray(value) else value
            if start % 256 == 0:
                log(f"  {name}: loaded {min(start + 16, len(columns)):,}/{len(columns):,} features")
    result.flush()
    return result


def prepare_shared(args):
    shared = args.shared.resolve()
    if getattr(args, 'rebuild_setup', False) and (shared / 'dataset_manifest.json').exists():
        archive_setup(args)
    if (shared / 'dataset_manifest.json').exists():
        manifest, _ = verify_shared(shared)
        if manifest['seed'] != args.seed or manifest['sample_per_level'] != args.sample_per_level or manifest['data_directory'] != manifest_path(args.data):
            raise ValueError('Shared setup differs; choose a new --shared directory.')
        log('Shared preparation already complete.')
        return
    for name in ('splits.csv', 'feature_definitions.json', 'feature_experiments.json'):
        if (shared / name).exists():
            raise ValueError('Incomplete shared preparation; inspect it or choose a new --shared directory.')
    if not args.metadata.is_file():
        raise FileNotFoundError(
            f'Learner metadata is missing: {args.metadata}. Restore the original workbook '
            'or use --metadata with a CSV containing text_id, learner_id, cefr_level. '
            'The feature dataframe alone cannot reconstruct learner-disjoint splits.')
    data = args.data.resolve()
    specs = dictionary(data)
    with duckdb.connect() as con:
        columns = [r[0] for r in con.execute('DESCRIBE SELECT * FROM read_parquet(?)', [str(data / 'feature_dataframe.parquet')]).fetchall()]
        if columns != ['text_id', 'cefr_level'] + [s['column'] for s in specs]:
            raise ValueError('Feature dictionary does not match dataframe column order.')
        texts = con.execute('SELECT text_id, cefr_level FROM read_parquet(?)', [str(data / 'feature_dataframe.parquet')]).fetchdf()
        levels = con.execute('SELECT text_id, cefr_level, level FROM read_parquet(?)', [str(data / 'text_metadata.parquet')]).fetchdf()
    if levels.text_id.duplicated().any() or levels.level.isna().any() or not levels.level.isin(range(1, 16)).all():
        raise ValueError('Course metadata must have unique IDs and levels 1–15.')
    learners = read_learner_metadata(args.metadata)
    frame = texts.merge(learners, on='text_id', validate='one_to_one', how='left', suffixes=('', '_metadata'))
    if frame.learner_id.isna().any() or not frame.cefr_level.eq(frame.cefr_level_metadata).all():
        raise ValueError('Learner metadata does not match dataframe.')
    frame = frame.drop(columns='cefr_level_metadata').merge(levels, on='text_id', validate='one_to_one', how='left', suffixes=('', '_metadata'))
    if frame.level.isna().any() or not frame.cefr_level.eq(frame.cefr_level_metadata).all():
        raise ValueError('Course levels do not match dataframe labels.')
    frame = frame.drop(columns='cefr_level_metadata')
    frame.level = frame.level.astype(int)
    if (frame.groupby('level').cefr_level.nunique() != 1).any():
        raise ValueError('Each course level must map to exactly one CEFR class.')
    if set(frame.level) != set(range(1, 16)):
        raise ValueError('Both prediction tasks require the full set of course levels 1–15.')
    frame['split'] = frame.learner_id.map(lambda learner: assign_split(learner, args.seed))
    if args.sample_per_level:
        frame = pd.concat([group.sample(min(args.sample_per_level, len(group)), random_state=args.seed)
                           for _, group in frame.groupby(['split', 'level'])], ignore_index=True)
    frame = frame.sort_values('text_id')
    labels = [label for label in LABELS if label in set(frame.cefr_level)]
    validate_splits(frame, labels)
    for role in ('train', 'validation', 'test'):
        counts = frame.loc[frame.split == role].level.value_counts()
        if any(counts.get(level, 0) < 2 for level in range(1, 16)):
            raise ValueError('Each course level needs at least two texts in each split.')
    complexity = [s['column'] for s in specs if s['group'] == 'complexity']
    accuracy = [s['column'] for s in specs if s['group'] == 'accuracy']
    all_features = complexity + accuracy
    if {v[k] for v in DERIVED.values() for k in ('numerator', 'denominator')} - set(all_features):
        raise ValueError('ERRANT counts needed for error proportions are missing.')
    experiments = {'majority': {'columns': [], 'rationale': 'Training majority baseline.'},
                   'complexity': {'columns': complexity, 'rationale': 'Complexity features.'},
                   'accuracy': {'columns': accuracy, 'rationale': 'Linguistic error features.'},
                   'combined': {'columns': all_features, 'rationale': 'Complexity and error features.'},
                   'combined_plus_error_composition': {'columns': all_features + list(DERIVED), 'rationale': 'Added relative error proportions.'}}
    shared.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.prepare_', dir=shared) as temporary:
        stage = Path(temporary)
        # writing_id is an alias for the profile-scoring and LLM preparation consumers.
        frame['writing_id'] = frame.text_id
        frame.to_csv(stage / 'splits.csv', index=False)
        write_json(stage / 'feature_definitions.json', {'original': specs, 'derived': DERIVED})
        write_json(stage / 'feature_experiments.json', experiments)
        manifest = {'schema_version': 2, 'seed': args.seed, 'sample_per_level': args.sample_per_level,
                    'scope': 'smoke_test' if args.sample_per_level else 'full_corpus',
                    'data_directory': manifest_path(data), 'labels': labels, 'rows': len(frame),
                    'level_to_cefr': {str(k): v for k, v in frame.groupby('level').cefr_level.first().items()},
                    'split_method': 'SHA256(seed:learner_id); 70/15/15; disjoint learners shared by all tasks',
                    'metadata_input': fingerprint(args.metadata),
                    'sources': [fingerprint(data / name) for name in ('feature_dataframe.parquet', 'feature_dictionary.csv', 'text_metadata.parquet')],
                    'files': {name: digest(stage / name) for name in ('splits.csv', 'feature_definitions.json', 'feature_experiments.json')},
                    'counts': frame.groupby(['split', 'level']).size().unstack(fill_value=0).to_dict(orient='index')}
        for name in manifest['files']:
            os.replace(stage / name, shared / name)
        write_json(shared / 'dataset_manifest.json', manifest)
    log(f'Prepared {len(frame):,} texts for both CEFR and level models.')


def verify_shared(shared):
    shared = Path(shared).resolve()
    manifest = read_json(shared / 'dataset_manifest.json')
    for source in manifest['sources']:
        if 'sha256' not in source:
            raise ValueError('This setup uses legacy timestamp checks. Use --rebuild-setup to '
                             'archive the old setup and results, then rebuild with content checksums.')
        if fingerprint(resolve_manifest_path(source['path'])) != source:
            raise ValueError(f"Input contents changed: {source['path']}. Use --rebuild-setup "
                             'to archive previous results and rebuild a consistent setup.')
    for name, expected in manifest['files'].items():
        if digest(shared / name) != expected:
            raise ValueError(f'Shared {name} changed; do not alter fixed splits or feature definitions.')
    frame = pd.read_csv(shared / 'splits.csv', dtype={'learner_id': str, 'cefr_level': str})
    validate_splits(frame, manifest['labels'])
    return manifest, frame


# Reviewed predecessor: prediction fitting is unchanged by the reporting/IG upgrade.
COMPATIBLE_CODE_HASHES = {
    '17bfa884110ca8f9a356e64385db625d1900ba2482a5a90222e7b07f67932793',
    '7b79daeb92cd2e70f50d780cff41776b44bece8dbda4ba775b82942a769c8bc0',
    'a49d3c5f83cc4e96153e6ca723232d0a103c69b6ede52bae34c4aee0318c7fb4',
    '56e6794518f4e3fc767adfab3e8d92049e5bc2f2bd9dd4cb6765fde13f05db61',
    '50477ce7ea47f9d9017c4afe7e5bec575867d37aea9fea12426ed99a08181346',
    'e0d62d154ffcce4cb4d3121bf7a38a53bd066b53dff4b1a2b5cb3a72b2562644',
    '54cc5adecba2753d94cea814fde7dbfbe4caaa31bb4734de84b1793b1751be13',
    'f6d730fe1072d755fc7b556430ec7d3649543ca353695a41d86f932c2b9b1786',
    'a903de7fb5a76e4864aab4ad7262e684bcb1b7f554f788739371b6a7e3561f59',
    '2b2d4a8a3ae7f875449e00626da6d547baf91b483715c705cca77c002a22e074',
    'a709ece9beb031e31575baaba337d0b90365aa05457c15f7dc1a46bd0c9d0e12',
    '0ef4d5d6a4fba161f119669203f85e6a6b8253040ee2b7889bd277975b4a1ead',
    # Majority report-label rename only; model fitting and rankings are unchanged.
    'd535c9d11846046597866ecb60281e858716fee3eb6f63906a44efc181babfbd',
    '2e80bbc4c33c6f77906e11c685f0628a47fccbd3ade660b1173ccb587489f32c',
    # Separate median naming; baseline estimators and feature calculations are unchanged.
    '1878e02bc73cfc6aae45e0b753ef279a3c94913119019bb800cf4b8f0ce546c8',
    '7e9d331a92fc49eb57954cd0026fd7d5663d36abe379816d5ed2709ef6f3de36',
    # Side-by-side workbook layout; ranking calculations and model inputs are unchanged.
    '6e93da35cda8409629fe5cd8ac42ac651bdaf22450e6a9e55a392f1ff78f1c43',
    '07ec471c1a2988226a43cbd85081f96123e607cfadb523076296054105ca0ef0',
    '70f3db5676b79b14d9d280e7b670ef06459cc4aa0f8bdfb2891477fe04f55a2d',
    # Export recovery and readable errors only; training and ranking calculations are unchanged.
    'e0499c0a13e8df4b1d7a6ceefce75661576117ab3523861af50da441d679ad5f',
    # POLKE experiment, MI rename and bootstrap standard deviations; model fitting is unchanged.
    '3816ff0e36229a9e90f78bdde968cbea3a25bda9f11f55816a6541719734ecd3',
    # Full standard-deviation column names; model fitting is unchanged.
    '4205e4b33da2a2652636639d4034ab3580acdfee6734e0781d155ec989473274',
    # Per-CEFR and per-level feature contribution tables; model fitting is unchanged.
    'ffbb4e2e03420b57d0898f5430bb912bf55d2d8aa142d1a1c2329946dbd9cdbb',
    '90b69f6bcf9d3c61fcdf9ae5d0f98522be76c4d7789d1ae93772bee2bb70aeb6',
    'fa10d647a8f8d0d5ecd91b6cc143c95cf239fb105a811947dde77ac19c01eaf8',
    # Short report names, per-model Per_CEFR/Per_Level folders and combined_without_text_size; fitting unchanged.
    '1f84c12bb934e86a99015312cad5b96ed83e797561a3f14a2a23294c71a57554',
    # Tool contribution tables; fitting unchanged.
    '4780a4f742a46c4e0232b50d19cec2cf8cee3969478346c24e5baf69edb645cb',
    # Without-length experiments from the section-1 length lists; fitting unchanged.
    '583bde81e0633e3b2214467c32e5a1f3e5b8350c473a31c5de39630eab4a1cfe',
    # Renamed length-free exports; new setups omit the retired experiment; fitting unchanged.
    '54d298cbd052a147e321b283328211b68d23c9a0c73a52213c098b304b3e2048',
    # List-1-only length experiments and the *_without_length_correlated_features rename; fitting unchanged.
    '1b9b9166eec5b975991e9f3cadd803fa51342c03b155826302e4355848f0843f',
    # Grouped report folders; fitting unchanged.
    '3c1aee97477a14596ab622e01285db61c1118e26eba0757eee975e079f57eb23',
    # Training lock for parallel training processes; fitting unchanged.
    '7265ca9606aaf6697c5d74ea88dfe9a243a172b850233ae7dae8c726bd02babf',
    # Duplicate-filtered experiments; fitting unchanged.
    '93b12e338e70d1ef2d31393cb21d4c6308fbf05c514a87bd9ee1a126598c198e',
}


MI_RATIONALE = 'Training-only Mutual Information subset.'
POLKE_RATIONALE = 'POLKE grammatical-construction features only, for comparison with POLKE-based studies.'
# Length-dependent features found in section 1: *_without_length_features excludes list 1 (the direct length features),
# *_without_length_correlated_features excludes both lists.
LENGTH_BLACKLISTS = (ROOT / '1_Complexity_Analysis_Full_Module/blacklist_1_direct_length_features.csv',
                     ROOT / '1_Complexity_Analysis_Full_Module/blacklist_2_length_correlated_features.csv')
WITHOUT_LENGTH_FEATURES_RATIONALE = ('{} features without the direct length features: the numbers of words and different words '
                                     '(section-1 length list 1).')
WITHOUT_LENGTH_CORRELATED_RATIONALE = ('{} features without any length-dependent feature: direct size measures, other raw counts '
                                       'and measures that follow text length (both section-1 length lists).')
LENGTH_ONLY_RATIONALE = ('Only the direct length features of section-1 length list 1 (numbers of words and different words): '
                         'how well text length alone predicts proficiency.')
# Near-duplicate features (training |Spearman| >= 0.95 with a kept feature); a kept feature is never on a length list
# that the same experiment removes, so each duplicate-filtered set keeps the information of what it drops.
DUPLICATE_LIST = ROOT / '1_Complexity_Analysis_Full_Module/blacklist_3_duplicate_features.csv'
DUPLICATE_EXPERIMENTS = (
    ('combined', 'combined_without_duplicate_features',
     'Combined features without duplicate features (section-1 duplicate list).'),
    ('combined_without_length_features', 'combined_without_length_and_duplicate_features',
     'Combined features without duplicate features and without the direct length features (length list 1).'),
    ('combined_without_length_correlated_features', 'combined_without_length_correlated_and_duplicate_features',
     'Combined features without duplicate features and without any length-dependent feature (length lists 1 and 2).'),
)
# Superseded by complexity_without_length_correlated_features but still listed in the fixed shared setup; its runs were
# archived outside the repository.
RETIRED_EXPERIMENTS = ('complexity_without_text_size',)
# Models exported for the LLM reference profiles, as (feature set, export variant).
SCORING_EXPORTS = (('complexity', 'all_features'),
                   ('complexity_without_length_correlated_features', 'without_length_correlated_features'))
# The retired complexity_without_text_size export and the former name of the length-free export; removed at the next export.
FORMER_SCORING_EXPORTS = ('without_direct_text_size_features', 'without_length_features')
# Renamed experiments as (task, former run folder, current run folder); training artifacts move unchanged.
RUN_RENAMES = (
    ('level', 'regressor__majority', 'regressor__median'),
    ('cefr', 'classifier__training_ranked_features', 'classifier__training_mutual_information_features'),
    ('level', 'classifier__training_ranked_features', 'classifier__training_mutual_information_features'),
    ('level', 'regressor__training_ranked_features', 'regressor__training_mutual_information_features'),
    ('cefr', 'classifier__complexity_without_length', 'classifier__complexity_without_length_correlated_features'),
    ('cefr', 'classifier__combined_without_length', 'classifier__combined_without_length_correlated_features'),
    ('level', 'classifier__complexity_without_length', 'classifier__complexity_without_length_correlated_features'),
    ('level', 'classifier__combined_without_length', 'classifier__combined_without_length_correlated_features'),
    ('level', 'regressor__complexity_without_length', 'regressor__complexity_without_length_correlated_features'),
    ('level', 'regressor__combined_without_length', 'regressor__combined_without_length_correlated_features'),
)


def canonical_training_configuration(configuration):
    """Read former experiment labels without rewriting training records."""
    if (configuration.get('task') == 'level' and configuration.get('algorithm') == 'regressor'
            and configuration.get('feature_set') == 'majority'
            and configuration.get('definition') == {'columns': [], 'rationale': 'Training majority baseline.'}):
        return {**configuration, 'feature_set': 'median',
                'definition': {'columns': [], 'rationale': 'Training median baseline.'}}
    definition = configuration.get('definition')
    if (configuration.get('feature_set') == 'training_ranked_features' and isinstance(definition, dict)
            and definition.get('rationale') == 'Optional training-only ranked feature subset.'):
        return {**configuration, 'feature_set': 'training_mutual_information_features',
                'definition': {**definition, 'rationale': MI_RATIONALE}}
    family = configuration.get('feature_set')
    if (family in ('complexity_without_length', 'combined_without_length') and isinstance(definition, dict)
            and definition.get('rationale') == WITHOUT_LENGTH_CORRELATED_RATIONALE.format(family.split('_')[0].capitalize())):
        return {**configuration, 'feature_set': f'{family}_correlated_features'}
    return dict(configuration)


def verify_run(run, configuration=None):
    marker = run / 'completed.json'
    if not marker.exists():
        return False
    saved = read_json(marker)
    if configuration is not None and token(configuration) != saved['configuration_id']:
        candidates = [run/'run_config.json', run/'ranking_config.json']
        previous = next((read_json(p) for p in candidates if p.is_file()), {})
        compatible = canonical_training_configuration(previous)
        compatible['code_sha256'] = configuration.get('code_sha256')
        if (previous.get('code_sha256') not in COMPATIBLE_CODE_HASHES
                or token(previous) != saved['configuration_id'] or compatible != configuration):
            raise ValueError(f'{run.name}: settings or training code changed. Choose a new --output directory.')
    for name, expected in saved['files'].items():
        if digest(run / name) != expected:
            raise ValueError(f'Changed run artifact: {run / name}')
    return True


def migrate_renamed_runs(folder, shared, task):
    """Move verified runs to their current experiment names; retain every training artifact byte."""
    root = Path(folder).resolve()
    parent = root / '1_Intermediate_Calculations'
    for renamed_task, old_name, new_name in RUN_RENAMES:
        if renamed_task != task:continue
        old, new = parent / old_name, parent / new_name
        if not old.exists() and not new.exists():continue
        for path in (old, new):
            if path.is_symlink() or not path.resolve().is_relative_to(root):
                raise ValueError(f'Run directory escapes its task folder: {path}')
        if old.exists() and new.exists():
            raise ValueError(f'Both {old_name} and {new_name} exist; resolve the duplicate before continuing.')
        run = old if old.exists() else new
        if not verify_run(run):raise ValueError(f'Cannot rename an incomplete run: {run}')
        config = canonical_training_configuration(read_json(run / 'run_config.json'))
        algorithm, family = new_name.split('__', 1)
        if (config.get('task'), config.get('algorithm'), config.get('feature_set')) != (task, algorithm, family):
            raise ValueError(f'Unexpected model in {run}.')
        seal_path = Path(shared) / 'final_evaluation.json'
        seal = read_json(seal_path) if seal_path.exists() else None
        update_seal = False
        if seal is not None:
            if task not in seal['locations'] or resolve_manifest_path(seal['locations'][task]) != root:
                raise ValueError(f'The sealed {task}-model location differs from the requested output.')
            frozen = seal['runs'][task]
            checksum = digest(run / 'completed.json')
            if old_name in frozen:
                if new_name in frozen or frozen[old_name] != checksum:
                    raise ValueError(f'The sealed {old_name} differs from its saved training record.')
                frozen[new_name] = frozen.pop(old_name)
                update_seal = True
            elif frozen.get(new_name) != checksum:
                raise ValueError(f'{new_name} is missing from the sealed training records.')
        if old.exists():
            old.rename(new)
            log(f'Renamed {task}/{old_name} to {new_name}; trained artifacts preserved.')
        # Also repairs an interrupted rename whose directory moved before the seal was saved.
        if update_seal:write_json(seal_path, seal)


def close_matrix(matrix):
    if isinstance(matrix, np.memmap) and not matrix._mmap.closed:
        matrix._mmap.close()


RANKING_SHEETS = {'mi': 'Mutual_Information', 'ig': 'Information_Gain'}
RANKING_SHEET = 'MI_And_IG'
RANKING_SHARED_COLUMNS = ['target', 'feature', 'eligible', 'training_observations']
RANKING_METHOD_COLUMNS = {
    'mi': {'mutual_information':'mi_score', 'rank':'mi_rank', 'selected':'mi_selected'},
    'ig': {'information_gain_bits':'ig_score_bits', 'rank':'ig_rank', 'selected':'ig_selected',
           'best_split_threshold':'ig_best_split_threshold'},
}
RANKING_METHODS = {
    'mi': 'training-only quantile-binned mutual information; 20 bins; missing separate',
    'ig': 'training-only best binary entropy split; minimum leaf 0.5%; finite observations only',
}


def combine_ranking_tables(tables):
    """One row per target and feature, retaining each method's own score and rank."""
    mi = tables['mi'].rename(columns=RANKING_METHOD_COLUMNS['mi'])
    ig = tables['ig'].rename(columns=RANKING_METHOD_COLUMNS['ig'])
    combined = mi.merge(ig, on=RANKING_SHARED_COLUMNS, how='outer', validate='one_to_one', indicator=True)
    if not combined['_merge'].eq('both').all() or combined.duplicated(['target','feature']).any():
        raise ValueError('MI and IG rankings must describe the same features and training observations.')
    combined = combined.drop(columns='_merge')
    combined['selected_by'] = np.select(
        [combined.mi_selected & combined.ig_selected, combined.mi_selected, combined.ig_selected],
        ['Both', 'MI only', 'IG only'], default='Neither')
    columns = RANKING_SHARED_COLUMNS + list(RANKING_METHOD_COLUMNS['mi'].values()) + list(RANKING_METHOD_COLUMNS['ig'].values()) + ['selected_by']
    return combined[columns].sort_values(['target','mi_rank','feature']).reset_index(drop=True)


def split_ranking_table(combined):
    """Restore independent method tables in rank order for existing model selection."""
    columns = RANKING_SHARED_COLUMNS + list(RANKING_METHOD_COLUMNS['mi'].values()) + list(RANKING_METHOD_COLUMNS['ig'].values()) + ['selected_by']
    if list(combined.columns) != columns or combined.duplicated(['target','feature']).any():
        raise ValueError('Invalid combined ranking table columns or duplicate features.')
    tables = {}
    for method, names in RANKING_METHOD_COLUMNS.items():
        table = combined[RANKING_SHARED_COLUMNS + list(names.values())].rename(columns={v:k for k,v in names.items()})
        original_columns = ['target','feature','eligible',next(iter(names)),'training_observations']
        if method == 'ig':original_columns.append('best_split_threshold')
        original_columns += ['rank','selected']
        tables[method] = table[original_columns].sort_values(['target','rank','feature']).reset_index(drop=True)
    expected = combine_ranking_tables(tables)
    actual = combined.sort_values(['target','mi_rank','feature']).reset_index(drop=True)
    if not expected.equals(actual):raise ValueError('Combined ranking table has inconsistent selection labels.')
    return tables


def ranking_settings(args, manifest, rows):
    return {'dataset_id': token(manifest), 'top_k': args.top_k,
            'min_observations': args.min_observations, 'rows': rows,
            'methods': RANKING_METHODS, 'code_sha256': digest(Path(__file__))}


def ranking_tables(folder, manifest):
    path = folder/'training_feature_ranking.xlsx'
    if not path.is_file():raise ValueError('Run training feature ranking first.')
    with pd.ExcelFile(path) as workbook:
        embedded = '_Pipeline_Metadata' in workbook.sheet_names
        if embedded:
            metadata = json.loads(pd.read_excel(workbook, sheet_name='_Pipeline_Metadata', header=None).iloc[0, 0])
            config = metadata['configuration']
            if token(config) != metadata['configuration_id']:
                raise ValueError('Ranking workbook configuration changed.')
        else:
            if not verify_run(folder):raise ValueError('Ranking workbook has no verified configuration.')
            config = read_json(folder/'ranking_config.json')
        combined = pd.read_excel(workbook, sheet_name=RANKING_SHEET) if RANKING_SHEET in workbook.sheet_names else None
        if combined is not None:
            if embedded and combined_ranking_digest(combined) != metadata.get('combined_table_sha256'):
                raise ValueError('Ranking workbook table changed; restore the verified workbook.')
            tables = split_ranking_table(combined)
        else:
            tables = {method: pd.read_excel(workbook, sheet_name=sheet) for method, sheet in RANKING_SHEETS.items()}
    if embedded and any(ranking_table_digest(table) != metadata['table_sha256'][method]
                        for method, table in tables.items()):
        raise ValueError('Ranking workbook table changed; restore the verified workbook.')
    if config.get('format') not in ('two_method_workbook', 'side_by_side_workbook'):
        raise ValueError('Run rank once to combine the saved MI and IG rankings.')
    if config['settings']['dataset_id'] != token(manifest):
        raise ValueError('Rankings belong to a different shared dataset.')
    if not embedded or combined is None:
        save_ranking_workbook(folder, config['settings'], tables, config['source_table_sha256'])
        return ranking_tables(folder, manifest)
    return config, tables


def ranking_table_digest(table):
    """Hash the values read from Excel, independent of workbook formatting."""
    return hashlib.sha256(table.to_json(orient='split', double_precision=15).encode('utf8')).hexdigest()


def combined_ranking_digest(table):
    """Allow worksheet sorting while still detecting changes to scores or selections."""
    if not {'target','feature'}.issubset(table.columns):raise ValueError('Invalid combined ranking table identifiers.')
    return ranking_table_digest(table.sort_values(['target','feature']).reset_index(drop=True))


def remove_legacy_ranking_tables(folder):
    # Remove only the superseded generated tables after the workbook is verified.
    paths = [folder/name for name in ('training_feature_ranking.csv', 'ranking_config.json', 'completed.json')] + [folder/'2_Information_Gain'/name
        for name in ('training_feature_ranking.csv', 'ranking_config.json', 'completed.json')]
    for path in paths:
        if path.is_file():
            if not path.resolve().is_relative_to(folder.resolve()):raise ValueError('Ranking path escapes its folder.')
            path.unlink()
    legacy = folder/'2_Information_Gain'
    if legacy.is_dir() and not any(legacy.iterdir()):legacy.rmdir()


def save_ranking_workbook(folder, settings, tables, source_hashes):
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    config = {'format':'side_by_side_workbook', 'settings':settings,
              'source_table_sha256':source_hashes,
              'source_note':'Original ranking table identities retained for existing experiment reproducibility.'}
    with tempfile.TemporaryDirectory(prefix='.ranking_export_',dir=folder) as temporary:
        stage=Path(temporary)
        with pd.ExcelWriter(stage/'training_feature_ranking.xlsx',engine='openpyxl') as writer:
            combined = combine_ranking_tables(tables)
            combined.to_excel(writer,sheet_name=RANKING_SHEET,index=False,freeze_panes=(1,2))
            sheet = writer.sheets[RANKING_SHEET]
            sheet.sheet_view.showGridLines = False
            sheet.row_dimensions[1].height = 42
            table = Table(displayName='TrainingFeatureRanking', ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            sheet.add_table(table)
            comments = {
                'mi_score':'Training-only mutual information (natural logarithm units). Higher means more dependence; compare ranks rather than averaging MI and IG scores.',
                'ig_score_bits':'Training-only information gain in bits from the best allowed binary threshold split.',
                'mi_selected':'Whether this feature is eligible and within the configured top-k MI ranks for this target.',
                'ig_selected':'Whether this feature is eligible and within the configured top-k IG ranks for this target.',
                'selected_by':'Both, MI only, IG only, or Neither. This comparison does not combine the model feature sets.',
            }
            for cell in sheet[1]:
                name = cell.value
                color = '167A70' if name.startswith('mi_') else '285E96' if name.startswith('ig_') else '334155'
                cell.fill = PatternFill('solid', fgColor=color)
                cell.font = Font(color='FFFFFF', bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical='center')
                sheet.column_dimensions[cell.column_letter].width = 64 if name=='feature' else 25 if name in ('training_observations','ig_best_split_threshold') else 17
                if name in comments:cell.comment=Comment(comments[name], 'Thesis pipeline')
            scores = {'mi_score','ig_score_bits','ig_best_split_threshold'}
            selection_colors = {'Both':'D1FAE5', 'MI only':'CCFBF1', 'IG only':'DBEAFE', 'Neither':'F1F5F9'}
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if combined.columns[cell.column-1] in scores:cell.number_format='0.000000'
                row[-1].fill=PatternFill('solid', fgColor=selection_colors[row[-1].value])
        # Keep provenance inside the workbook; retain historical selection IDs so
        # existing fitted models can still be reused without changing their inputs.
        workbook_path = stage/'training_feature_ranking.xlsx'
        saved_combined = pd.read_excel(workbook_path, sheet_name=RANKING_SHEET)
        saved_tables = split_ranking_table(saved_combined)
        hashes = {method: ranking_table_digest(table) for method, table in saved_tables.items()}
        for method in tables:
            expected = tables[method].sort_values(['target','rank','feature']).reset_index(drop=True)
            pd.testing.assert_frame_equal(expected, saved_tables[method], check_exact=False, rtol=1e-14, atol=0)
        metadata = {'configuration':config, 'configuration_id':token(config), 'table_sha256':hashes,
                    'combined_table_sha256':combined_ranking_digest(saved_combined)}
        with pd.ExcelWriter(workbook_path, engine='openpyxl', mode='a') as writer:
            pd.DataFrame([[json.dumps(metadata, sort_keys=True)]]).to_excel(
                writer, sheet_name='_Pipeline_Metadata', index=False, header=False)
            writer.sheets['_Pipeline_Metadata'].sheet_state = 'hidden'
        if combined_ranking_digest(pd.read_excel(workbook_path, sheet_name=RANKING_SHEET)) != metadata['combined_table_sha256']:
            raise ValueError('Ranking workbook export did not preserve its table.')
        try:
            os.replace(workbook_path,folder/'training_feature_ranking.xlsx')
        except PermissionError as exc:
            raise PermissionError('Close training_feature_ranking.xlsx in Excel, then rerun the same command.') from exc
    remove_legacy_ranking_tables(folder)


def rank_features(manifest, training, columns, min_observations, top_k, working, memory):
    """MI and IG of every feature with CEFR and course level, from the given training rows only;
    the cross-validation ranks each fold's training part the same way."""
    rows={'mi':[],'ig':[]}
    for start in range(0,len(columns),16):
        batch=columns[start:start+16]
        matrix=load_matrix(manifest,training,batch,working,'ranking',memory)
        try:
            for index,name in enumerate(batch):
                values=np.asarray(matrix[:,index]);finite=np.isfinite(values)
                usable=finite.sum()>=min_observations and np.unique(values[finite]).size>1
                bins=np.zeros(len(values),dtype=int)
                if usable:
                    edges=np.unique(np.quantile(values[finite],np.linspace(0,1,21)))
                    bins[finite]=np.searchsorted(edges[1:-1],values[finite],side='right')+1
                for target,labels in [('cefr',training.cefr_level),('level',training.level)]:
                    identity={'target':target,'feature':name,'eligible':bool(usable)}
                    observations=int(finite.sum())
                    rows['mi'].append({**identity,'mutual_information':float(mutual_info_score(labels,bins)) if usable else 0.,
                                       'training_observations':observations})
                    gain=0.;threshold=None
                    if usable:
                        model=DecisionTreeClassifier(criterion='entropy',max_depth=1,
                            min_samples_leaf=max(2,int(np.ceil(observations*.005))),random_state=manifest['seed'])
                        model.fit(values[finite,None],np.asarray(labels)[finite]);tree=model.tree_
                        if tree.children_left[0]>=0:
                            left,right=tree.children_left[0],tree.children_right[0]
                            gain=max(0.,float(tree.impurity[0]-(tree.weighted_n_node_samples[left]*tree.impurity[left]+
                                tree.weighted_n_node_samples[right]*tree.impurity[right])/tree.weighted_n_node_samples[0]))
                            threshold=float(tree.threshold[0])
                    rows['ig'].append({**identity,'information_gain_bits':gain,'training_observations':observations,
                                       'best_split_threshold':threshold})
        finally:close_matrix(matrix)
        log(f'Ranked {min(start+16,len(columns))}/{len(columns)} features with MI and IG using training texts only.')
    tables={}
    for method,score in [('mi','mutual_information'),('ig','information_gain_bits')]:
        table=pd.DataFrame(rows[method]).sort_values(['target','eligible',score,'feature'],ascending=[True,False,False,True])
        table['rank']=table.groupby('target').cumcount()+1
        table['selected']=table.eligible & table['rank'].le(top_k)
        tables[method]=table
    return tables


def feature_ranking(args):
    manifest, frame = verify_shared(args.shared)
    training = frame[frame.split == 'train'].sort_values('text_id')
    settings = ranking_settings(args, manifest, len(training))
    args.ranking.mkdir(parents=True, exist_ok=True)
    if (args.ranking/'training_feature_ranking.xlsx').exists():
        config,_=ranking_tables(args.ranking,manifest)
        saved=dict(config['settings']);old_hash=saved.pop('code_sha256')
        desired=dict(settings);desired.pop('code_sha256')
        if saved!=desired or old_hash not in COMPATIBLE_CODE_HASHES | {settings['code_sha256']}:
            raise ValueError('Ranking settings or calculation code changed; use a new ranking/output setup.')
        remove_legacy_ranking_tables(args.ranking)
        log('Matching MI and IG rankings already exist in training_feature_ranking.xlsx.')
        return
    # Import previous completed rankings without changing feature selections or model identities.
    legacy={'mi':args.ranking,'ig':args.ranking/'2_Information_Gain'}
    if all((folder/'completed.json').exists() for folder in legacy.values()):
        tables={}; hashes={}
        for method,folder in legacy.items():
            verify_run(folder)
            previous=read_json(folder/'ranking_config.json')
            for key in ('dataset_id','top_k','min_observations','rows'):
                if previous[key]!=settings[key]:raise ValueError('Saved ranking settings differ from this run.')
            if previous['method']!=RANKING_METHODS[method]:raise ValueError('Unexpected saved ranking method.')
            if previous['code_sha256'] not in COMPATIBLE_CODE_HASHES | {settings['code_sha256']}:
                raise ValueError('Saved ranking code is not a reviewed compatible version.')
            path=folder/'training_feature_ranking.csv'
            hashes[method]=digest(path)
            tables[method]=pd.read_csv(path,float_precision='round_trip')
        save_ranking_workbook(args.ranking,settings,tables,hashes)
        log('Combined the existing MI and IG rankings into one workbook; feature selections retained.')
        return
    if (args.shared/'final_evaluation.json').exists():
        raise ValueError('Shared evaluation is sealed; rankings cannot change.')
    if (args.ranking/'completed.json').exists():
        raise ValueError('Only one legacy ranking is complete. Use a new ranking folder to calculate both methods.')
    definitions=read_json(args.shared/'feature_definitions.json')
    columns=sorted({s['column'] for s in definitions['original']} | set(DERIVED))
    with tempfile.TemporaryDirectory(prefix='.rank_',dir=args.ranking) as temporary:
        tables=rank_features(manifest,training,columns,args.min_observations,args.top_k,Path(temporary),args.memory)
    hashes={method:hashlib.sha256(table.to_csv(index=False).encode('utf8')).hexdigest() for method,table in tables.items()}
    save_ranking_workbook(args.ranking,settings,tables,hashes)


def experiment_definitions(args, manifest):
    definitions=read_json(args.shared/'feature_experiments.json')
    for name in RETIRED_EXPERIMENTS:
        definitions.pop(name,None)
    if args.task == 'level':
        definitions['median'] = {'columns': [], 'rationale': 'Training median baseline.'}
    # Derived from the fixed feature definitions and the length lists, so the shared setup stays unchanged.
    specs=read_json(args.shared/'feature_definitions.json')['original']
    if all(Path(path).is_file() for path in LENGTH_BLACKLISTS):
        direct,correlated=(set(pd.read_csv(path).column) for path in LENGTH_BLACKLISTS)
        for suffix,excluded,rationale in (('without_length_features',direct,WITHOUT_LENGTH_FEATURES_RATIONALE),
                                          ('without_length_correlated_features',direct|correlated,WITHOUT_LENGTH_CORRELATED_RATIONALE)):
            for family in ('complexity','combined'):
                definitions[f'{family}_{suffix}']={
                    'columns':[column for column in definitions[family]['columns'] if column not in excluded],
                    'rationale':rationale.format(family.capitalize())}
        if DUPLICATE_LIST.is_file():
            duplicates=set(pd.read_csv(DUPLICATE_LIST).column)
            for base,name,rationale in DUPLICATE_EXPERIMENTS:
                definitions[name]={'columns':[column for column in definitions[base]['columns'] if column not in duplicates],
                                   'rationale':rationale}
        definitions['length_only']={'columns':[spec['column'] for spec in specs if spec['column'] in direct],
                                    'rationale':LENGTH_ONLY_RATIONALE}
    polke=[spec['column'] for spec in specs if spec['source']=='POLKE']
    if polke:
        definitions['polke']={'columns':polke,'rationale':POLKE_RATIONALE}
    if args.use_ranked:
        config,tables=ranking_tables(args.ranking,manifest)
        for method,name,rationale in [
            ('mi','training_mutual_information_features',MI_RATIONALE),
            ('ig','training_information_gain_features','Training-only best-threshold Information Gain subset.')]:
            table=tables[method]
            selected=table[(table.target==args.task) & table.selected].feature.tolist()
            if not selected:raise ValueError(f'{method.upper()} ranking has no eligible selected features.')
            definitions[name]={'columns':selected,'rationale':rationale,
                               'ranking_sha256':config['source_table_sha256'][method]}
    return definitions


def class_metrics(truth, predicted, labels):
    order = {value: index for index, value in enumerate(labels)}
    a = np.array([order[x] for x in truth]); b = np.array([order[x] for x in predicted])
    kappa = cohen_kappa_score(a, b, labels=list(range(len(labels))), weights='quadratic')
    return {'rows': len(a), 'accuracy': float(accuracy_score(truth, predicted)),
            'macro_f1': float(f1_score(truth, predicted, labels=labels, average='macro', zero_division=0)),
            'ordinal_mae': float(np.abs(a-b).mean()), 'within_one_accuracy': float((np.abs(a-b) <= 1).mean()),
            'quadratic_weighted_kappa': float(kappa) if np.isfinite(kappa) else None,
            'labels': labels, 'confusion_matrix': confusion_matrix(truth, predicted, labels=labels).tolist(),
            'per_class': classification_report(truth, predicted, labels=labels, output_dict=True, zero_division=0)}


def evaluate_model(model, matrix, rows, manifest, task, algorithm):
    table = rows[['text_id', 'learner_id', 'level', 'cefr_level']].copy()
    labels = manifest['labels']
    mapping = {int(k): v for k, v in manifest['level_to_cefr'].items()}
    if task == 'cefr':
        probabilities = model.predict_proba(matrix)
        table['predicted_cefr'] = model.classes_[probabilities.argmax(axis=1)]
        for i, label in enumerate(model.classes_):
            table[f'probability_cefr_{label}'] = probabilities[:, i]
        return {'cefr': class_metrics(table.cefr_level, table.predicted_cefr, labels)}, table
    if algorithm == 'classifier':
        probabilities = model.predict_proba(matrix)
        table['predicted_level'] = model.classes_[probabilities.argmax(axis=1)].astype(int)
        grouped = np.zeros((len(rows), len(labels)))
        for i, level in enumerate(model.classes_):
            table[f'probability_level_{int(level)}'] = probabilities[:, i]
            grouped[:, labels.index(mapping[int(level)])] += probabilities[:, i]
        for i, label in enumerate(labels):
            table[f'probability_cefr_{label}'] = grouped[:, i]
        table['predicted_cefr'] = np.asarray(labels)[grouped.argmax(axis=1)]
        regression = None
    else:
        continuous = np.clip(model.predict(matrix), 1., 15.)
        table['predicted_level_continuous'] = continuous
        table['predicted_level'] = np.rint(continuous).astype(int)
        table['predicted_cefr'] = table.predicted_level.map(mapping)
        regression = {'mae': float(mean_absolute_error(table.level, continuous)),
                      'rmse': float(np.sqrt(mean_squared_error(table.level, continuous))),
                      'r_squared': float(r2_score(table.level, continuous)),
                      'spearman_correlation': float(spearmanr(table.level, continuous).statistic) if np.unique(continuous).size>1 and np.unique(table.level).size>1 else None}
    metrics = {'cefr': class_metrics(table.cefr_level, table.predicted_cefr, labels),
               'level': class_metrics(table.level, table.predicted_level, list(range(1, 16)))}
    if regression is not None:
        metrics['regression'] = regression
    return metrics, table


def model_contributions(model, matrix, columns, task, algorithm):
    labels = list(model.classes_) if algorithm == 'classifier' else ['numerical_level']
    totals = np.zeros((len(columns), len(labels)))
    for start in range(0, len(matrix), 512):
        values = np.abs(model.eval_terms(matrix[start:start+512]))
        if values.ndim == 2:
            values = values[:, :, None]
        totals += values.sum(axis=0)
    return pd.DataFrame([{'feature': column, 'output': label, 'mean_absolute_contribution': totals[i,j]/len(matrix)}
                         for i,column in enumerate(columns) for j,label in enumerate(labels)])


def training_configuration(args, manifest, family, algorithm, definition):
    return {'dataset_id': token(manifest), 'task': args.task, 'algorithm': algorithm,
            'feature_set': family, 'definition': definition, 'rounds': args.rounds,
            'min_observations': args.min_observations, 'seed': manifest['seed'],
            'interactions': 0, 'max_bins': 128, 'learning_rate': .04, 'outer_bags': 1,
            'class_weight': 'balanced' if algorithm == 'classifier' and family != 'majority' else 'none',
            'code_sha256': digest(Path(__file__)),
            'packages': {name: importlib.metadata.version(name) for name in ('interpret','scikit-learn','numpy','duckdb','pandas','joblib')}}


@contextmanager
def training_lock(run, poll_seconds=60):
    """Let parallel training processes share runs: the first one trains a run, and another process
    that reaches the same run waits for it, then reuses the result."""
    lock = run.parent / f'.{run.name}.lock'
    waiting = False
    while True:
        try:
            descriptor = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            break
        except FileExistsError:
            try:
                content, age = lock.read_text(encoding='ascii', errors='replace'), time.time() - lock.stat().st_mtime
            except OSError:
                continue  # released meanwhile
            owner = int(content) if content.strip().isdigit() else 0
            if owner == 0 and age < 60:
                time.sleep(1)  # the owner has not written its process ID yet
                continue
            # A lock left by a process that stopped, or by this process after an error, is stale.
            if owner in (0, os.getpid()) or not psutil.pid_exists(owner):
                lock.unlink(missing_ok=True)
                continue
            if not waiting:
                log(f'{run.name}: being trained by process {owner}; waiting for it')
                waiting = True
            time.sleep(poll_seconds)
    try:
        with os.fdopen(descriptor, 'w', encoding='ascii') as handle:
            handle.write(str(os.getpid()))
        yield
    finally:
        lock.unlink(missing_ok=True)


def fit_run(manifest, training, validation, columns, task, algorithm, family, rounds, min_observations,
            workers, memory, work, label, score=None):
    """Fit one model on the training rows exactly as every pipeline run is fitted, and evaluate it on the
    held-out rows. score(model, matrix, features) runs on the held-out matrix before it is closed.
    The cross-validation fits each fold with this function."""
    train_matrix=valid_matrix=None
    try:
        selected=[]; excluded=[]
        target = training.cefr_level if task == 'cefr' else training.level
        if family in ('majority', 'median'):
            train_matrix=np.zeros((len(training),1));valid_matrix=np.zeros((len(validation),1))
            model = DummyClassifier(strategy='most_frequent') if algorithm == 'classifier' else DummyRegressor(strategy='median')
            model.fit(train_matrix,target)
        else:
            train_matrix=load_matrix(manifest,training,columns,work,'training',memory)
            keep=[]
            for i,column in enumerate(columns):
                finite=train_matrix[:,i][np.isfinite(train_matrix[:,i])]
                reason='insufficient_training_observations' if len(finite)<min_observations else 'constant_in_training' if np.ptp(finite)==0 else ''
                if reason:excluded.append({'feature':column,'reason':reason})
                else:keep.append(i)
            selected=[columns[i] for i in keep]
            if not selected:raise ValueError('No usable training features.')
            filtered=np.lib.format.open_memmap(work/'filtered.npy',mode='w+',dtype=np.float64,shape=(len(training),len(keep)))
            for start in range(0,len(training),4096):filtered[start:start+4096]=train_matrix[start:start+4096,keep]
            close_matrix(train_matrix);train_matrix=filtered
            valid_matrix=load_matrix(manifest,validation,selected,work,'validation',memory)
            estimator=ExplainableBoostingClassifier if algorithm=='classifier' else ExplainableBoostingRegressor
            model=estimator(feature_names=selected,feature_types=['continuous']*len(selected),max_bins=128,interactions=0,
                            validation_size=0,outer_bags=1,learning_rate=.04,max_rounds=rounds,
                            smoothing_rounds=min(50,rounds//4),early_stopping_rounds=0,n_jobs=workers,random_state=manifest['seed'])
            log(f'{label}: fitting {len(training):,} texts and {len(selected):,} features')
            weights=compute_sample_weight('balanced',target) if algorithm=='classifier' else None
            model.fit(train_matrix,target,sample_weight=weights)
        metrics,predictions=evaluate_model(model,valid_matrix,validation,manifest,task,algorithm)
        scored=score(model,valid_matrix,selected) if score is not None and selected else None
        return model,selected,excluded,metrics,predictions,scored
    finally:close_matrix(train_matrix);close_matrix(valid_matrix)


def train_models(args):
    manifest, frame = verify_shared(args.shared)
    if (args.shared / 'final_evaluation.json').exists():
        raise ValueError('Both tasks are sealed after final testing starts. Use a new shared experiment.')
    definitions = experiment_definitions(args, manifest)
    requested = list(args.experiments or definitions)
    # Existing commands selecting the baseline family still include its regression counterpart.
    if args.task == 'level' and 'majority' in requested and 'median' not in requested:
        requested.append('median')
    training = frame[frame.split == 'train'].sort_values('text_id')
    validation = frame[frame.split == 'validation'].sort_values('text_id')
    originals = {s['column'] for s in read_json(args.shared / 'feature_definitions.json')['original']}
    algorithms = ['classifier'] if args.task == 'cefr' else (args.algorithms or ['classifier','regressor'])
    for family in requested:
        if family not in definitions or not re.fullmatch(r'[a-z0-9_]+', family):
            raise ValueError(f'Unknown feature set {family}')
        columns = definitions[family]['columns']
        if len(columns) != len(set(columns)) or set(columns) - (originals | set(DERIVED)):
            raise ValueError('Metadata or undeclared feature in experiment.')
        for algorithm in algorithms:
            if (family == 'majority' and algorithm != 'classifier') or (family == 'median' and algorithm != 'regressor'):
                continue
            name = f'{algorithm}__{family}'
            config = training_configuration(args, manifest, family, algorithm, definitions[family])
            run = args.output / '1_Intermediate_Calculations' / name
            if verify_run(run, config):
                log(f'{args.task}/{name}: complete; reusing')
                continue
            run.parent.mkdir(parents=True, exist_ok=True)
            with training_lock(run):
                # A parallel training process may have finished this run while this one waited.
                if verify_run(run, config):
                    log(f'{args.task}/{name}: trained by a parallel process; reusing')
                    continue
                if run.exists():
                    raise ValueError(f'Incomplete run exists: {run}; inspect it or use a new output.')
                with tempfile.TemporaryDirectory(prefix='.fit_', dir=run.parent) as temporary:
                    work=Path(temporary); stage=work/'result';stage.mkdir()
                    model,selected,excluded,metrics,predictions,contributions=fit_run(
                        manifest,training,validation,columns,args.task,algorithm,family,args.rounds,args.min_observations,
                        args.workers,args.memory,work,f'{args.task}/{name}',
                        lambda model,matrix,features:model_contributions(model,matrix,features,args.task,algorithm))
                    write_json(stage/'run_config.json',config);write_json(stage/'feature_columns.json',selected)
                    write_json(stage/'excluded_features.json',excluded);write_json(stage/'validation_metrics.json',metrics)
                    predictions.to_csv(stage/'validation_predictions.csv',index=False)
                    if contributions is not None:contributions.to_csv(stage/'feature_contributions.csv',index=False)
                    joblib.dump({'model':model,'features':selected,'dataset_id':token(manifest),'task':args.task,'algorithm':algorithm},stage/'model.joblib',compress=3)
                    write_json(stage/'completed.json',{'configuration_id':token(config),'files':{p.name:digest(p) for p in stage.iterdir() if p.is_file()}})
                    stage.rename(run)
                    log(f"Validation CEFR macro F1: {metrics['cefr']['macro_f1']:.4f}")


def final_test(args):
    manifest,frame=verify_shared(args.shared)
    # Freeze every currently completed run across both tasks before reading any test labels.
    locations={'cefr':args.cefr_output.resolve(),'level':args.level_output.resolve()}
    snapshots={}
    for task,folder in locations.items():
        snapshots[task]={run.name:digest(run/'completed.json') for run in (folder/'1_Intermediate_Calculations').glob('*') if run.is_dir() and verify_run(run)}
        if not snapshots[task]:raise ValueError('Complete both CEFR and level experiments before final testing.')
        for name in snapshots[task]:
            config=read_json(folder/'1_Intermediate_Calculations'/name/'run_config.json')
            if config['dataset_id']!=token(manifest) or config['task']!=task:
                raise ValueError('Final-test tasks do not share the same dataset and target definitions.')
    if not args.runs or any(name not in snapshots[args.task] for name in args.runs):
        raise ValueError('Select completed run names using --runs.')
    lock={'dataset_id':token(manifest),'locations':{k:manifest_path(v) for k,v in locations.items()},'runs':snapshots}
    lock_path=args.shared/'final_evaluation.json'
    if lock_path.exists():
        previous=read_json(lock_path)
        previous['locations']={key:manifest_path(resolve_manifest_path(value)) for key,value in previous['locations'].items()}
        if previous!=lock:raise ValueError('Final evaluation was sealed with different runs or locations.')
    write_json(lock_path,lock)
    rows=frame[frame.split=='test'].sort_values('text_id')
    for name in args.runs:
        run=locations[args.task]/'1_Intermediate_Calculations'/name
        marker=run/'test_completed.json'
        if marker.exists():
            if not all(digest(run/file)==checksum for file,checksum in read_json(marker).items()):raise ValueError('Changed test output.')
            continue
        bundle=joblib.load(run/'model.joblib')
        if bundle['dataset_id']!=token(manifest):raise ValueError('Different dataset in model bundle.')
        with tempfile.TemporaryDirectory(prefix='.test_',dir=run.parent) as temporary:
            matrix=None;work=Path(temporary)
            try:
                matrix=load_matrix(manifest,rows,bundle['features'],work,'test',args.memory) if bundle['features'] else np.zeros((len(rows),1))
                metrics,predictions=evaluate_model(bundle['model'],matrix,rows,manifest,args.task,bundle['algorithm'])
                write_json(work/'test_metrics.json',metrics);predictions.to_csv(work/'test_predictions.csv',index=False)
                for filename in ('test_metrics.json','test_predictions.csv'):os.replace(work/filename,run/filename)
                write_json(marker,{filename:digest(run/filename) for filename in ('test_metrics.json','test_predictions.csv')})
            finally:close_matrix(matrix)

def course_level_metrics(predictions, levels, labels, experiment, split):
    table = predictions.merge(levels, on='text_id', how='left', validate='one_to_one', suffixes=('', '_metadata'))
    if table.level.isna().any() or not table.cefr_level.eq(table.cefr_level_metadata).all():
        raise ValueError('Prediction and course metadata do not match.')
    order = {label: index for index, label in enumerate(labels)}
    if not table.predicted_cefr.isin(labels).all():
        raise ValueError('Unknown predicted CEFR label.')
    table['distance'] = (table.cefr_level.map(order) - table.predicted_cefr.map(order)).abs()
    return [{'experiment': experiment, 'split': split, 'level': int(level),
             'cefr_level': group.cefr_level.iloc[0], 'support': len(group),
             'cefr_accuracy': float(group.distance.eq(0).mean()),
             'cefr_ordinal_mae': float(group.distance.mean()),
             'within_one_cefr_accuracy': float(group.distance.le(1).mean())}
            for level, group in table.groupby('level')]


def archive_setup(args):
    """Preserve generated artifacts before an explicitly requested rebuild."""
    from datetime import datetime
    if not args.metadata.is_file():
        raise FileNotFoundError(f'Restore learner metadata before rebuilding: {args.metadata}')
    read_learner_metadata(args.metadata)
    plans = []
    for folder, names in [
        (args.shared, ['dataset_manifest.json', 'splits.csv', 'feature_definitions.json', 'feature_experiments.json', 'final_evaluation.json']),
        (args.ranking, ['completed.json', 'ranking_config.json', 'training_feature_ranking.csv', 'training_feature_ranking.xlsx', '2_Information_Gain']),
        (args.cefr_output, ['1_Experiments', '2_Reports', '1_Intermediate_Calculations', '2_Prediction_Performance', '3_Feature_Contributions']),
        (args.level_output, ['1_Experiments', '2_Reports', '1_Intermediate_Calculations', '2_Prediction_Performance', '3_Feature_Contributions']),
        (args.llm_output/'2_Level_Reference_Profiles', ['1_Intermediate_Calculations', '2_Profile_Validation', '3_Reusable_Models_For_LLMs']),
        (args.llm_output/'1_CEFR_Reference_Profiles', ['1_Intermediate_Calculations', '2_Profile_Validation', '3_Reusable_Models_For_LLMs']),
        (args.baseline_output, ['1_Leaderboards','2_Model_Comparisons','2_Feature_Set_Comparisons',
            'cefr_leaderboard.csv','level_leaderboard.csv','baseline_improvements.csv','feature_selection_comparison.csv']),
        (args.llm_output, ['post_training_resources.csv'])]:
        for name in names:
            source = folder / name
            if source.exists():
                if source.is_symlink() or not source.resolve().is_relative_to(BASE.resolve()):
                    raise ValueError('Automatic archival supports only generated outputs inside this EBM folder. Use new custom output folders.')
                plans.append(source)
    destination = BASE / '9_Previous_Runs' / datetime.now().strftime('setup_%Y%m%d_%H%M%S_%f')
    for source in plans:
        target = destination / source.relative_to(BASE)
        target.parent.mkdir(parents=True, exist_ok=True)
        source.rename(target)
    log(f'Archived previous setup and outputs: {destination}')


def organize_saved_outputs():
    """Migrate generated reports only; immutable training artifacts stay intact."""
    def checked(path):
        if not path.resolve().is_relative_to(BASE.resolve()):
            raise ValueError(f'Output path escapes the EBM folder: {path}')
        return path

    def move_folder(source, target):
        if not source.exists():return
        checked(source);checked(target)
        target.mkdir(parents=True, exist_ok=True)
        for source_file in source.rglob('*'):
            if not source_file.is_file():continue
            target_file=target/source_file.relative_to(source)
            checked(source_file);checked(target_file)
            target_file.parent.mkdir(parents=True, exist_ok=True)
            if target_file.exists():
                if digest(source_file) != digest(target_file):
                    raise ValueError(f'Two different reports exist: {source_file} and {target_file}')
                source_file.unlink()
            else:source_file.rename(target_file)
        for folder in sorted((p for p in source.rglob('*') if p.is_dir()),key=lambda p:len(p.parts),reverse=True):
            checked(folder).rmdir()
        source.rmdir()

    for task, name, destination in (
        ('2_CEFR_Classification','3_CEFR_Profile_Weights','1_CEFR_Reference_Profiles'),
        ('3_Level_Prediction','3_Level_Profile_Weights','2_Level_Reference_Profiles')):
        move_folder(BASE/task/name, BASE/'5_Data_For_LLM_Post_Training'/destination)
    for task in ('2_CEFR_Classification','3_Level_Prediction'):
        root=BASE/task
        move_folder(root/'2_Analysis_Reports',root/'2_Prediction_Performance')
        reports=root/'2_Prediction_Performance'
        for path in reports.rglob('*.xlsx'):checked(path).unlink()
        old=reports/'1_Performance_Metrics/direct_and_level_derived_cefr_comparison.csv'
        if old.exists():
            destination=BASE/'2_CEFR_Classification/1_Intermediate_Calculations/cefr_prediction_routes_comparison.csv'
            if not destination.exists():
                destination.parent.mkdir(parents=True,exist_ok=True)
                old.rename(checked(destination))
            else:checked(old).unlink()
        for name in ('1_CEFR_Reference_Profiles','2_Level_Reference_Profiles'):
            profile=BASE/'5_Data_For_LLM_Post_Training'/name
            if not profile.exists():continue
            move_folder(profile/'2_Analysis_Reports',profile/'2_Profile_Validation')
            for path in profile.rglob('*.xlsx'):checked(path).unlink()
            old_features=profile/'1_Intermediate_Calculations/2_Normalization/scoring_features.csv'
            if old_features.exists():
                new_features=old_features.with_name('feature_weights.csv')
                if not new_features.exists():checked(old_features).rename(checked(new_features))
            for path in (profile/'3_Reusable_Models_For_LLMs').glob('scorer_*.json'):
                destination=path.with_name(path.name.replace('scorer_','weights_',1))
                if destination.exists():
                    if digest(path)!=digest(destination):raise ValueError('Conflicting old and new profile weights.')
                    checked(path).unlink()
                else:checked(path).rename(checked(destination))


def per_label_report_rows(metrics):
    """Accuracy within an actual label's texts; equal to recall when support > 0.

    This is not one-versus-rest accuracy, which also counts true negatives.
    Leave accuracy undefined for labels with no actual texts.
    """
    for index, label in enumerate(metrics['labels']):
        counts = metrics['confusion_matrix'][index]
        support = sum(counts)
        yield {'label':label, 'accuracy':counts[index]/support if support else np.nan,
               **metrics['per_class'][str(label)]}


BOOTSTRAP_RESAMPLES = 2000
BOOTSTRAP_SEED = 42
BOOTSTRAP_CHUNK = 100
CLASS_METRICS = ('accuracy', 'macro_f1', 'ordinal_mae', 'within_one_accuracy', 'quadratic_weighted_kappa')
REGRESSION_METRICS = ('mae', 'rmse', 'r_squared', 'spearman_correlation')
LABEL_METRICS = ('accuracy', 'precision', 'recall', 'f1-score')
COURSE_LEVEL_METRICS = ('cefr_accuracy', 'cefr_ordinal_mae', 'within_one_cefr_accuracy')
_LEARNER_DRAWS = {}
_RUN_RESAMPLES = {}


def learner_draws(frame, split):
    """How often each learner is drawn in every bootstrap resample of one split.

    A drawn learner brings all of their texts, because texts by one learner are not independent.
    Every model evaluated on the split sees the same draws, so model differences are paired.
    """
    rows = frame[frame.split == split].sort_values('text_id')
    key = (split, hashlib.sha256(pd.util.hash_pandas_object(rows[['text_id', 'learner_id']].astype(str), index=False).to_numpy().tobytes()).hexdigest())
    if key not in _LEARNER_DRAWS:
        codes, learners = pd.factorize(rows.learner_id.astype(str), sort=True)
        count = len(learners)
        rng = np.random.default_rng(BOOTSTRAP_SEED)
        draws = np.empty((BOOTSTRAP_RESAMPLES, count), dtype=np.uint8)
        for start in range(0, BOOTSTRAP_RESAMPLES, BOOTSTRAP_CHUNK):
            size = min(BOOTSTRAP_CHUNK, BOOTSTRAP_RESAMPLES - start)
            picks = rng.integers(0, count, size=(size, count)) + np.arange(size)[:, None] * count
            tally = np.bincount(picks.ravel(), minlength=size * count).reshape(size, count)
            if tally.max() > np.iinfo(np.uint8).max:raise ValueError('Too many repeated learners in one resample.')
            draws[start:start + size] = tally
        _LEARNER_DRAWS[key] = (rows.text_id.to_numpy(), codes, draws)
    return _LEARNER_DRAWS[key]


def confusion_resamples(confusion):
    """class_metrics for a stack of confusion matrices (resample x actual x predicted)."""
    index = np.arange(confusion.shape[1])
    distance = np.abs(index[:, None] - index[None, :])
    total = confusion.sum(axis=(1, 2))
    actual, predicted = confusion.sum(axis=2), confusion.sum(axis=1)
    hits = confusion[:, index, index]
    with np.errstate(divide='ignore', invalid='ignore'):
        recall = hits / actual
        expected = actual[:, :, None] * predicted[:, None, :] / total[:, None, None]
        chance = (distance ** 2 * expected).sum(axis=(1, 2))
        kappa = np.where(chance > 0, 1 - (distance ** 2 * confusion).sum(axis=(1, 2)) / chance, np.nan)
        f1 = np.where(actual + predicted > 0, 2 * hits / (actual + predicted), 0.)
        return {'accuracy': hits.sum(axis=1) / total, 'macro_f1': f1.mean(axis=1),
                'ordinal_mae': (distance * confusion).sum(axis=(1, 2)) / total,
                'within_one_accuracy': confusion[:, distance <= 1].sum(axis=1) / total,
                'quadratic_weighted_kappa': kappa,
                # Zero-division conventions follow classification_report and per_label_report_rows.
                'labels': {'accuracy': np.where(actual > 0, recall, np.nan),
                           'precision': np.where(predicted > 0, hits / predicted, 0.),
                           'recall': np.where(actual > 0, recall, 0.), 'f1-score': f1}}


def weighted_correlation(weights, first, second):
    """Pearson correlation of every weighted resample (weights: resample x text)."""
    total = weights.sum(axis=1, keepdims=True)
    first = first - (weights * first).sum(axis=1, keepdims=True) / total
    second = second - (weights * second).sum(axis=1, keepdims=True) / total
    variance = (weights * first ** 2).sum(axis=1) * (weights * second ** 2).sum(axis=1)
    with np.errstate(divide='ignore', invalid='ignore'):
        return np.where(variance > 0, (weights * first * second).sum(axis=1) / np.sqrt(variance), np.nan)


def weighted_ranks(weights, groups, inverse):
    """Average rank of each text inside every weighted resample; ties share ranks as in spearmanr."""
    counts = np.asarray((groups.T @ weights.T).T)
    return (np.cumsum(counts, axis=1) - counts + (counts + 1) / 2)[:, inverse]


def prepare_resampling(predictions, manifest, task):
    """Sparse indicators that turn per-text resample weights into confusion matrices and sums."""
    rows = np.arange(len(predictions))
    order = {label: index for index, label in enumerate(manifest['labels'])}
    pairs = {'cefr': (predictions.cefr_level.map(order), predictions.predicted_cefr.map(order), len(order))}
    if task == 'level':
        pairs['level'] = (predictions.level - 1, predictions.predicted_level - 1, 15)
    prepared = {'targets': {}}
    for target, (actual, predicted, size) in pairs.items():
        if actual.isna().any() or predicted.isna().any() or not actual.between(0, size - 1).all() or not predicted.between(0, size - 1).all():
            raise ValueError(f'Unknown {target} label in saved predictions.')
        codes = actual.to_numpy(int) * size + predicted.to_numpy(int)
        prepared['targets'][target] = (sparse.csr_matrix((np.ones(len(codes)), (rows, codes)), shape=(len(codes), size * size)), size)
    if 'predicted_level_continuous' in predictions:
        truth, estimate = predictions.level.to_numpy(float), predictions.predicted_level_continuous.to_numpy(float)
        prepared['regression'] = {'truth': truth, 'estimate': estimate, 'ranks': None}
        if np.unique(truth).size > 1 and np.unique(estimate).size > 1:
            ranks = []
            for values in (truth, estimate):
                unique, inverse = np.unique(values, return_inverse=True)
                ranks.append((sparse.csr_matrix((np.ones(len(values)), (rows, inverse)), shape=(len(values), len(unique))), inverse))
            prepared['regression']['ranks'] = ranks
    if task == 'cefr':
        levels, inverse = np.unique(predictions.level.to_numpy(int), return_inverse=True)
        distance = np.abs(pairs['cefr'][0].to_numpy(float) - pairs['cefr'][1].to_numpy(float))
        prepared['course_levels'] = levels
        prepared['course'] = {name: sparse.csr_matrix((values, (rows, inverse)), shape=(len(rows), len(levels)))
                              for name, values in (('texts', np.ones(len(rows))), ('cefr_accuracy', (distance == 0) * 1.),
                                                   ('cefr_ordinal_mae', distance), ('within_one_cefr_accuracy', (distance <= 1) * 1.))}
    return prepared


def resample_metrics(weights, prepared):
    """Every reported metric for each row of per-text weights (resample x text)."""
    result = {}
    for target, (indicator, size) in prepared['targets'].items():
        result[target] = confusion_resamples(np.asarray((indicator.T @ weights.T).T).reshape(-1, size, size))
    if 'regression' in prepared:
        truth, estimate = prepared['regression']['truth'], prepared['regression']['estimate']
        total = weights.sum(axis=1)
        squared = weights @ (estimate - truth) ** 2
        mean = weights @ truth / total
        spread_of_truth = weights @ truth ** 2 - total * mean ** 2
        with np.errstate(divide='ignore', invalid='ignore'):
            r_squared = np.where(spread_of_truth > 0, 1 - squared / spread_of_truth, np.nan)
        ranks = prepared['regression']['ranks']
        result['regression'] = {'mae': weights @ np.abs(estimate - truth) / total, 'rmse': np.sqrt(squared / total),
                                'r_squared': r_squared,
                                'spearman_correlation': np.full(len(weights), np.nan) if ranks is None else
                                weighted_correlation(weights, *(weighted_ranks(weights, *pair) for pair in ranks))}
    if 'course' in prepared:
        sums = {name: np.asarray((matrix.T @ weights.T).T) for name, matrix in prepared['course'].items()}
        with np.errstate(divide='ignore', invalid='ignore'):
            result['course_level'] = {name: sums[name] / sums['texts'] for name in COURSE_LEVEL_METRICS}
    return result


def run_resamples(run, split, frame, manifest, task):
    """Bootstrap distribution of each metric reported for a saved run (one value per resample)."""
    path = run / f'{split}_predictions.csv'
    key = (str(path.resolve()), digest(path))
    if key in _RUN_RESAMPLES:return _RUN_RESAMPLES[key]
    text_ids, learners, draws = learner_draws(frame, split)
    predictions = pd.read_csv(path).sort_values('text_id', kind='stable').reset_index(drop=True)
    if not np.array_equal(predictions.text_id.to_numpy(), text_ids):
        raise ValueError(f'{run.name}: {split} predictions do not match the {split} texts.')
    prepared = prepare_resampling(predictions, manifest, task)
    # The same calculation on the original texts must reproduce the saved metrics.
    saved = read_json(run / f'{split}_metrics.json')
    original = resample_metrics(np.ones((1, len(predictions))), prepared)
    checks = [(saved[target][name], original[target][name][0]) for target in prepared['targets'] for name in CLASS_METRICS]
    checks += [(saved['regression'][name], original['regression'][name][0]) for name in REGRESSION_METRICS if 'regression' in saved]
    for expected, value in checks:
        if (expected is None) != (not np.isfinite(value)) or (expected is not None and not np.isclose(expected, value, rtol=1e-9, atol=1e-12)):
            raise ValueError(f'{run.name}: bootstrap metrics do not reproduce the saved {split} metrics.')
    parts = [resample_metrics(draws[start:start + BOOTSTRAP_CHUNK][:, learners].astype(np.float64), prepared)
             for start in range(0, BOOTSTRAP_RESAMPLES, BOOTSTRAP_CHUNK)]

    def join(pieces):
        return {name: join([piece[name] for piece in pieces]) for name in pieces[0]} if isinstance(pieces[0], dict) else np.concatenate(pieces)
    result = join(parts)
    if 'course_levels' in prepared:result['course_levels'] = prepared['course_levels']
    _RUN_RESAMPLES[key] = result
    return result


def spread(values):
    """Bootstrap standard deviation, i.e. the metric's standard error; None if undefined in over 5% of resamples."""
    values = np.asarray(values, dtype=float)
    finite = values[np.isfinite(values)]
    return float(finite.std(ddof=1)) if finite.size >= max(2, .95 * values.size) else None


def with_spread(values, resamples, names, column=None):
    """Copy a row, placing each named metric's bootstrap standard deviation right after it."""
    row = {}
    for key, value in values.items():
        row[key] = value
        if key in names:
            samples = resamples[key] if column is None else resamples[key][:, column]
            row[f'{key}_standard_deviation'] = None if value is None or (isinstance(value, float) and np.isnan(value)) else spread(samples)
    return row


def collect_reports(folder, manifest, task, frame):
    comparisons=[];per_cefr=[];per_level=[];matrices=[];effects=[]
    for run in sorted((folder/'1_Intermediate_Calculations').glob('*')):
        if not run.is_dir() or not verify_run(run):continue
        config=canonical_training_configuration(read_json(run/'run_config.json'))
        if config['dataset_id']!=token(manifest) or config['task']!=task:raise ValueError('Report model uses a different dataset or task.')
        for split in ('validation','test'):
            if not (run/f'{split}_metrics.json').exists():continue
            if split=='test':
                if not (run/'test_completed.json').exists():continue
                if not all(digest(run/f)==h for f,h in read_json(run/'test_completed.json').items()):raise ValueError('Test report artifacts changed.')
            metrics=read_json(run/f'{split}_metrics.json'); predictions=pd.read_csv(run/f'{split}_predictions.csv')
            resamples=run_resamples(run,split,frame,manifest,task)
            info={'experiment':run.name,'task':task,'algorithm':config['algorithm'],'feature_set':config['feature_set'],'split':split,'rounds':config['rounds'],'class_weight':config['class_weight']}
            for target in ('cefr','level'):
                if target not in metrics:continue
                values=metrics[target]
                comparisons.append({**info,'evaluated_target':target,'features':len(read_json(run/'feature_columns.json')),'rows':values['rows'],
                                    **with_spread({k:values[k] for k in CLASS_METRICS},resamples[target],CLASS_METRICS),
                                    **{f'regression_{k}':v for k,v in with_spread(metrics.get('regression',{}),resamples.get('regression',{}),REGRESSION_METRICS).items()}})
                for i,details in enumerate(per_label_report_rows(values)):
                    label=details['label']
                    row={**info,**with_spread(details,resamples[target]['labels'],LABEL_METRICS,column=i)}
                    (per_cefr if target=='cefr' else per_level).append(row)
                    for j,predicted in enumerate(values['labels']):
                        matrices.append({**info,'evaluated_target':target,'actual':label,'predicted':predicted,'texts':values['confusion_matrix'][i][j]})
            if task=='cefr':
                course=course_level_metrics(predictions.drop(columns='level'),predictions[['text_id','cefr_level','level']],manifest['labels'],run.name,split)
                columns={int(level):index for index,level in enumerate(resamples['course_levels'])}
                per_level.extend({**info,**with_spread(row,resamples['course_level'],COURSE_LEVEL_METRICS,column=columns[row['level']]),
                                  'evaluation':'CEFR accuracy within actual course level; not level prediction'} for row in course)
        if (run/'feature_contributions.csv').exists():
            effects.extend({'experiment':run.name,**row} for row in pd.read_csv(run/'feature_contributions.csv').to_dict('records'))
    return comparisons,per_cefr,per_level,matrices,effects


def classification_percentages(table):
    """Percentage of each actual class assigned to each predicted class (0-100)."""
    table = table.copy()
    groups = ['experiment', 'task', 'split', 'evaluated_target', 'actual']
    totals = table.groupby(groups, dropna=False)['texts'].transform('sum')
    table['percentage'] = table['texts'].div(totals.where(totals.gt(0))).mul(100).fillna(0.0)
    return table


def save_report_tables(destination, tables):
    """Stage generated CSV reports and replace only files whose contents changed."""
    destination = Path(destination).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    for name in tables:
        if not (destination / name).resolve().is_relative_to(destination):
            raise ValueError('Report path escapes its output folder.')
    # Staged beside the destination, so the deepest staged paths stay as short as the final ones.
    with tempfile.TemporaryDirectory(prefix='.reports_', dir=destination.parent) as temporary:
        stage = Path(temporary)
        for name, table in tables.items():
            path = stage / name
            path.parent.mkdir(parents=True, exist_ok=True)
            table.to_csv(path, index=False)
        for name in tables:
            source, target = stage / name, destination / name
            target.parent.mkdir(parents=True, exist_ok=True)
            if target.exists() and digest(target) == digest(source):
                continue
            try:
                os.replace(source, target)
            except PermissionError as exc:
                raise PermissionError(f'Close {target.name} in Excel or any other program, then rerun the same command.') from exc


# Report folders group the models; level reports first split classifiers and regressors. Inside a group,
# a model folder drops what the group path already says, which keeps paths under Windows' 260 characters.
REPORT_GROUPS = (
    ('1_Top_Performing', {'complexity': 'complexity', 'combined': 'combined',
                          'combined_plus_error_composition': 'combined_plus_error_composition'}),
    ('2_Length_Features_Filtered/1_Without_Length_Features',
     {'complexity_without_length_features': 'complexity', 'combined_without_length_features': 'combined'}),
    ('2_Length_Features_Filtered/2_Without_Length_Correlated_Features',
     {'complexity_without_length_correlated_features': 'complexity', 'combined_without_length_correlated_features': 'combined'}),
    ('3_Mutual_Information_And_Information_Gain',
     {'training_mutual_information_features': 'mutual_information', 'training_information_gain_features': 'information_gain'}),
    ('4_Only_POLKE', {'polke': 'polke'}),
    ('5_Simple_Models', {'accuracy': 'accuracy', 'majority': 'majority', 'median': 'median', 'length_only': 'length_only'}),
    ('6_Duplicate_Features_Filtered/1_Duplicates', {'combined_without_duplicate_features': 'combined'}),
    ('6_Duplicate_Features_Filtered/2_Duplicates_And_Length', {'combined_without_length_and_duplicate_features': 'combined'}),
    ('6_Duplicate_Features_Filtered/3_Duplicates_And_Length_Correlated',
     {'combined_without_length_correlated_and_duplicate_features': 'combined'}),
)
OTHER_MODELS = '7_Other_Models'
ALGORITHM_FOLDERS = {'classifier': '1_Classifier', 'regressor': '2_Regressor'}
REPORT_GROUP_FOLDERS = ({part for group, _ in REPORT_GROUPS for part in group.split('/')}
                        | {OTHER_MODELS, *ALGORITHM_FOLDERS.values()})


def ranked_experiment_folders(quality, experiments, task):
    """Each experiment's report folder, numbered by validation rank: 2_Length_Features_Filtered/
    1_Without_Length_Features/07_combined for the CEFR task, 1_Classifier/... or 2_Regressor/... for the level task."""
    names = sorted(set(experiments))
    if any(not re.fullmatch(r'(classifier|regressor)__[a-z0-9_]+', name) for name in names):
        raise ValueError('Invalid experiment folder name.')
    ranks = [quality.get(name) for name in names]
    if any(rank is None or not np.isfinite(rank) or rank < 1 or int(rank) != rank for rank in ranks) or len(set(ranks)) != len(ranks):
        raise ValueError('Every experiment needs a unique positive validation rank.')
    folders = {}
    for name in names:
        algorithm, family = name.split('__', 1)
        group, short = next(((group, shorts[family]) for group, shorts in REPORT_GROUPS if family in shorts), (OTHER_MODELS, family))
        top = f'{ALGORITHM_FOLDERS[algorithm]}/' if task == 'level' else ''
        folders[name] = f'{top}{group}/{int(quality[name]):02d}_{short}'
    return folders


def remove_old_report_files(root, paths):
    root = Path(root).resolve()
    for path in paths:
        if not path.resolve().is_relative_to(root):
            raise ValueError('Old report path escapes the output folder.')
        if path.is_file():
            try:
                path.unlink()
            except PermissionError as exc:
                raise PermissionError(f'Close {path.name} in Excel or any other program, then rerun the same command.') from exc
    for folder in sorted({p.parent for p in paths}, key=lambda p: len(p.parts), reverse=True):
        if folder != root and folder.is_dir() and not any(folder.iterdir()):
            folder.rmdir()


def grouped_report_name(name, folder, task):
    """performance_metrics.csv becomes performance_metrics_cefr_05.csv in 1_Top_Performing/05_combined and
    performance_metrics_cefr_all_experiments.csv in 0_All_Experiments.

    The folder already names the model; repeating it pushed long paths past Excel's 218-character limit.
    The task keeps CEFR and level copies apart, so Excel can open any two reports at once.
    """
    path = Path(name)
    suffix = 'all_experiments' if folder == '0_All_Experiments' else Path(folder).name.split('_', 1)[0]
    return f'{path.stem}_{task}_{suffix}{path.suffix}'


def former_report_name(name, folder):
    """The earlier naming, e.g. performance_metrics_05_classifier__combined.csv; kept for cleanup."""
    path = Path(name)
    return f'{path.stem}_{folder}{path.suffix}'


def write_grouped_reports(destination, tables, quality, experiments, task, obsolete_names=()):
    """Keep folder numbers synchronized with validation ranks after every report refresh."""
    folders = ranked_experiment_folders(quality, experiments, task)
    outputs = {f'0_All_Experiments/{grouped_report_name(name, "0_All_Experiments", task)}': table
               for name, table in tables.items()}
    for experiment, folder in folders.items():
        for name, table in tables.items():
            outputs[f'{folder}/{grouped_report_name(name, folder, task)}'] = table[table.experiment.eq(experiment)]
    save_report_tables(destination, outputs)
    obsolete = []
    # Model folders at any depth: the grouped layout, the earlier flat one and folders left by rank changes.
    for folder in [destination/'0_All_Experiments', *model_folders(destination),
                   *(path for path in destination.iterdir() if path.is_dir() and re.fullmatch(r'(?:classifier|regressor)__[a-z0-9_]+', path.name))]:
        if not folder.is_dir():continue
        relative = folder.relative_to(destination).as_posix()
        for name in set(tables) | set(obsolete_names):
            for filename in (name, former_report_name(name, folder.name), grouped_report_name(name, folder.name, task)):
                if f'{relative}/{filename}' not in outputs:
                    obsolete.append(folder/filename)
    remove_old_report_files(destination, obsolete)
    remove_empty_report_folders(destination)


# One table per classifier output inside the model's numbered folder, named like the other reports:
# 1_Top_Performing/01_complexity/Per_CEFR/feature_contributions_cefr_01_A1.csv or
# 1_Classifier/1_Top_Performing/01_combined/Per_Level/feature_contributions_level_01_12.csv.
OUTPUT_TABLE_NAME = re.compile(r'feature_contributions_(?:cefr|level)_\d+_(?:[ABC][12]|[0-9]{1,2})\.csv')
OUTPUT_TABLE_FOLDERS = ('Per_CEFR', 'Per_Level')
# The previous layout collected the tables in one Per_CEFR or Per_Level folder per task, named by model.
FLAT_OUTPUT_TABLE_NAME = re.compile(r'feature_contributions_classifier__[a-z0-9_]+_(?:[ABC][12]|[0-9]{1,2})\.csv')
# Grouped model folders (01_combined) and the earlier flat ones (01_classifier__combined); group folders are capitalized.
MODEL_FOLDER_NAME = re.compile(r'\d+_(?:(?:classifier|regressor)__)?[a-z0-9_]+')


def model_folders(destination):
    """Numbered model folders at any depth of a report section, outside temporary staging folders."""
    destination = Path(destination)
    return [path for path in destination.rglob('*') if path.is_dir() and MODEL_FOLDER_NAME.fullmatch(path.name)
            and not any(part.startswith('.') for part in path.relative_to(destination).parts)]


def remove_empty_report_folders(destination):
    """Drop the model, per-output and group folders that a rank change or the regrouping left empty."""
    destination = Path(destination)
    folders = [path for path in destination.rglob('*') if path.is_dir()
               and not any(part.startswith('.') for part in path.relative_to(destination).parts)
               and (path.name in REPORT_GROUP_FOLDERS or path.name in OUTPUT_TABLE_FOLDERS or MODEL_FOLDER_NAME.fullmatch(path.name))]
    for folder in sorted(folders, key=lambda path: len(path.parts), reverse=True):
        if not any(folder.iterdir()):folder.rmdir()


def write_output_contribution_tables(destination, contributions, folders, task):
    """For each CEFR band (CEFR classifiers) or course level (level classifiers): every feature's
    mean absolute contribution to that output's score, most relevant first.

    Regressors have one numerical output and baselines have no features, so neither gets these tables.
    """
    outputs = {}
    for experiment, folder in folders.items():
        rows = contributions[contributions.experiment.eq(experiment)]
        if not experiment.startswith('classifier__') or rows.empty:continue
        for output, group in rows.groupby('output', sort=False):
            name = f'feature_contributions_{task}_{Path(folder).name.split("_", 1)[0]}_{output}.csv'
            if not OUTPUT_TABLE_NAME.fullmatch(name):raise ValueError(f'Unexpected model output {output!r} in {experiment}.')
            subfolder = 'Per_CEFR' if str(output) in LABELS else 'Per_Level'
            outputs[f'{folder}/{subfolder}/{name}'] = (
                group[['feature', 'mean_absolute_contribution']]
                .rename(columns={'feature': 'feature_name', 'mean_absolute_contribution': 'contribution'})
                .sort_values(['contribution', 'feature_name'], ascending=[False, True]))
    save_report_tables(destination, outputs)
    models = model_folders(destination)
    stale = [path for folder in models for subfolder in OUTPUT_TABLE_FOLDERS if (folder / subfolder).is_dir()
             for path in (folder / subfolder).iterdir()
             if path.is_file() and OUTPUT_TABLE_NAME.fullmatch(path.name)
             and f'{folder.relative_to(destination).as_posix()}/{subfolder}/{path.name}' not in outputs]
    stale += [path for subfolder in OUTPUT_TABLE_FOLDERS if (destination / subfolder).is_dir()
              for path in (destination / subfolder).iterdir() if path.is_file() and FLAT_OUTPUT_TABLE_NAME.fullmatch(path.name)]
    remove_old_report_files(destination, stale)
    # A model folder left behind by a rank change may now hold nothing but emptied subfolders.
    remove_empty_report_folders(destination)


TOOL_CONTRIBUTION_METHOD = "mean over validation texts of the absolute sum of the tool's feature contributions"


def feature_and_tool_contributions(model, matrix, features, algorithm):
    """Mean absolute contribution of each feature and each tool to each output over the rows of matrix.
    A tool counts as one combined feature: its features' contributions are summed for each text first."""
    outputs = [str(label) for label in model.classes_] if algorithm == 'classifier' else ['numerical_level']
    tools = {}
    for index, name in enumerate(features):
        tools.setdefault(name.split('__')[1], []).append(index)
    per_feature = np.zeros((len(features), len(outputs)))
    per_tool = {tool: np.zeros(len(outputs)) for tool in tools}
    for start in range(0, len(matrix), 512):
        values = model.eval_terms(matrix[start:start + 512])
        if values.ndim == 2:
            values = values[:, :, None]
        per_feature += np.abs(values).sum(axis=0)
        for tool, indices in tools.items():
            per_tool[tool] += np.abs(values[:, indices].sum(axis=1)).sum(axis=0)
    return outputs, tools, per_feature / len(matrix), {tool: value / len(matrix) for tool, value in per_tool.items()}


def score_tool_contributions(run, matrix, features):
    """Re-score a saved model on the validation matrix and group its feature contributions by tool."""
    bundle = joblib.load(run / 'model.joblib')
    if bundle['features'] != list(features):
        raise ValueError(f'{run.name}: the saved model and its feature list differ.')
    outputs, tools, per_feature, per_tool = feature_and_tool_contributions(bundle['model'], matrix, features, bundle['algorithm'])
    saved = pd.read_csv(run / 'feature_contributions.csv', dtype={'output': str})
    expected = saved.pivot(index='feature', columns='output', values='mean_absolute_contribution').loc[list(features), outputs]
    if not np.allclose(per_feature, expected.to_numpy(), rtol=1e-9, atol=1e-12):
        raise ValueError(f'{run.name}: re-scoring does not reproduce the saved feature contributions.')
    return [{'tool': tool, 'output': output, 'features': len(indices), 'contribution': float(per_tool[tool][position])}
            for tool, indices in tools.items() for position, output in enumerate(outputs)]


def tool_contribution_rows(folder, manifest, frame, memory='auto'):
    """Per model and output, how strongly each tool's features together move the score on validation texts.

    A tool counts as one combined feature: its features' contributions are summed for each text before the
    absolute value is taken, so features that offset each other are not double-counted. Re-scoring must
    reproduce the saved feature contributions; results are cached in the run folder while the run is unchanged.
    """
    results, pending = {}, {}
    for run in sorted((folder / '1_Intermediate_Calculations').glob('*')):
        if not run.is_dir() or not verify_run(run) or not read_json(run / 'feature_columns.json'):continue
        cache = run / 'tool_contributions.json'
        stamp = {'source_completed_sha256': digest(run / 'completed.json'), 'method': TOOL_CONTRIBUTION_METHOD}
        saved = read_json(cache) if cache.is_file() else {}
        if all(saved.get(key) == value for key, value in stamp.items()):
            results[run.name] = saved['rows']
        else:
            pending.setdefault(tuple(read_json(run / 'feature_columns.json')), []).append((run, cache, stamp))
    rows = frame[frame.split == 'validation'].sort_values('text_id')
    for features, members in pending.items():
        with tempfile.TemporaryDirectory(prefix='.reports_', dir=folder / '1_Intermediate_Calculations') as temporary:
            matrix = load_matrix(manifest, rows, list(features), Path(temporary), 'validation', memory)
            try:
                for run, cache, stamp in members:
                    log(f'{run.name}: grouping feature contributions by tool on {len(rows):,} validation texts')
                    results[run.name] = score_tool_contributions(run, matrix, features)
                    write_json(cache, {**stamp, 'validation_texts': len(rows), 'rows': results[run.name]})
            finally:
                close_matrix(matrix)
    return results


def tool_contribution_table(tool_rows, quality):
    """Long table: one row per model, output (plus 'overall', the mean over outputs) and tool."""
    table = pd.DataFrame([{'experiment': name, **row} for name, rows in tool_rows.items() for row in rows],
                         columns=['experiment', 'tool', 'output', 'features', 'contribution'])
    overall = table.groupby(['experiment', 'tool'], as_index=False).agg(
        features=('features', 'first'), contribution=('contribution', 'mean')).assign(output='overall')
    table = pd.concat([overall, table], ignore_index=True)
    table['share'] = table.contribution / table.groupby(['experiment', 'output']).contribution.transform('sum')
    position = table.output.map(lambda output: -1 if output == 'overall' else LABELS.index(output) if output in LABELS
                                else int(output) if str(output).isdigit() else 0)
    table = table.assign(position=position).sort_values(['experiment', 'position', 'contribution', 'tool'],
                                                        ascending=[True, True, False, True])
    table['tool_rank'] = table.groupby(['experiment', 'output']).cumcount() + 1
    table.insert(0, 'model_validation_rank', table.experiment.map(quality))
    table = table.sort_values(['model_validation_rank', 'experiment', 'position', 'tool_rank'], kind='stable')
    return table[['model_validation_rank', 'experiment', 'output', 'tool_rank', 'tool', 'features', 'contribution', 'share']]


def write_contribution_reports(folder, effects, quality, experiments, task, tool_rows=None):
    """Publish separate model explanations plus a combined, experiment-labelled view."""
    folder = Path(folder).resolve()
    experiments = sorted(set(experiments))
    if any(not re.fullmatch(r'(classifier|regressor)__[a-z0-9_]+', name) for name in experiments):
        raise ValueError('Invalid experiment folder name in contribution reports.')
    contributions = pd.DataFrame(effects, columns=[
        'experiment', 'feature', 'output', 'mean_absolute_contribution'])
    if not set(contributions.experiment).issubset(experiments):
        raise ValueError('Contribution data has no matching completed experiment.')
    ranking = contributions.groupby(['experiment', 'feature'], as_index=False).mean_absolute_contribution.mean()
    ranking = ranking.sort_values(
        ['experiment', 'mean_absolute_contribution', 'feature'], ascending=[True, False, True])
    ranking['overall_feature_rank'] = ranking.groupby('experiment').cumcount() + 1
    ranking = ranking.rename(columns={'mean_absolute_contribution':'overall_mean_absolute_contribution'})
    contributions = contributions.merge(ranking, on=['experiment','feature'], how='left', validate='many_to_one')
    contributions.insert(0, 'model_validation_rank', contributions.experiment.map(quality))
    contributions = contributions.sort_values(
        ['model_validation_rank', 'experiment', 'overall_feature_rank', 'output'], kind='stable')
    # Majority and median baselines retain headers but no invented linguistic weights.
    write_grouped_reports(folder / '3_Feature_Contributions', {
        'feature_contributions.csv': contributions,
        'tool_contributions.csv': tool_contribution_table(tool_rows or {}, quality)}, quality, experiments, task,
        obsolete_names=('post_model_feature_ranking.csv',))
    write_output_contribution_tables(folder / '3_Feature_Contributions', contributions,
                                     ranked_experiment_folders(quality, experiments, task), task)
    # Remove only the superseded summaries after every replacement has been saved.
    legacy = folder / '2_Prediction_Performance/2_Feature_Contributions'
    if not legacy.resolve().is_relative_to(folder):
        raise ValueError('Legacy contribution folder escapes the task output folder.')
    for name in ('feature_contributions.csv', 'post_model_feature_ranking.csv'):
        path = legacy / name
        if not path.resolve().is_relative_to(folder):
            raise ValueError('Legacy contribution file escapes the task output folder.')
        if path.is_file():
            path.unlink()
    if legacy.is_dir() and not any(legacy.iterdir()):
        legacy.rmdir()


def write_reports(args):
    manifest,frame=verify_shared(args.shared)
    comparisons,per_cefr,per_level,matrices,effects=collect_reports(args.output,manifest,args.task,frame)
    if not comparisons:raise ValueError('No completed experiments to report.')
    destination=args.output/'2_Prediction_Performance';destination.mkdir(parents=True,exist_ok=True)
    comparison=pd.DataFrame(comparisons).sort_values(['evaluated_target','split','macro_f1','accuracy','ordinal_mae','experiment'],ascending=[True,True,False,False,True,True])
    comparison.insert(0, 'quality_rank', comparison.groupby(['evaluated_target','split']).cumcount()+1)
    outputs={'performance_metrics.csv':comparison,
             'per_cefr_metrics.csv':pd.DataFrame(per_cefr),
             'per_level_metrics.csv':pd.DataFrame(per_level),
             ('CEFR_classifications.csv' if args.task=='cefr' else 'level_classifications.csv'):classification_percentages(pd.DataFrame(matrices))}
    primary = comparison[comparison.evaluated_target.eq(args.task) & comparison.split.eq('validation')]
    quality = primary.set_index('experiment').quality_rank.to_dict()
    for name, table in outputs.items():
        if 'experiment' not in table or 'quality_rank' in table:continue
        table.insert(0, 'model_validation_rank', table.experiment.map(quality))
        order = ['model_validation_rank', 'experiment']
        if 'rank' in table:order.append('rank')
        outputs[name] = table.sort_values(order, kind='stable')
    all_cefr=[]
    for task,folder in [('cefr',args.cefr_output),('level',args.level_output)]:
        if (folder/'1_Intermediate_Calculations').exists():
            all_cefr.extend(row for row in collect_reports(folder,manifest,task,frame)[0] if row['evaluated_target']=='cefr')
    if all_cefr:
        routes=pd.DataFrame(all_cefr).sort_values(['split','macro_f1','accuracy','ordinal_mae','experiment'],ascending=[True,False,False,True,True])
        routes.insert(0, 'quality_rank', routes.groupby('split').cumcount()+1)
        routes['prediction_route'] = np.where(routes.task.eq('cefr'), 'Direct CEFR classification',
            np.where(routes.algorithm.eq('classifier'), 'Level probabilities summed into CEFR bands', 'Numerical level rounded and mapped to CEFR'))
        route_path=args.cefr_output/'1_Intermediate_Calculations/cefr_prediction_routes_comparison.csv'
        route_path.parent.mkdir(parents=True,exist_ok=True)
        routes.to_csv(route_path,index=False)
    write_grouped_reports(destination, outputs, quality, comparison.experiment.unique(), args.task,
                          obsolete_names=('performance_metrics_comparison.csv','feature_set_comparison.csv'))
    write_contribution_reports(args.output, effects, quality, comparison.experiment.unique(), args.task,
                               tool_contribution_rows(args.output, manifest, frame, args.memory))
    remove_old_report_files(destination, [destination/'1_Performance_Metrics'/name for name in
        ('feature_set_comparison.csv','per_cefr_metrics.csv','per_level_metrics.csv')] +
        [destination/name for name in ('classifications.csv','CEFR_classifications.csv','level_classifications.csv')] +
        [destination/'0_All_Experiments/feature_set_comparison.csv'])
    export_scoring_inputs(args,manifest)
    log(f'Saved {args.task} reports: {destination}')


def export_scoring_inputs(args,manifest):
    definitions={s['column']:s for s in read_json(args.shared/'feature_definitions.json')['original']}
    profile_folder = '1_CEFR_Reference_Profiles' if args.task == 'cefr' else '2_Level_Reference_Profiles'
    destination=args.llm_output/profile_folder/'3_Reusable_Models_For_LLMs'
    algorithm = 'classifier' if args.task == 'cefr' else 'regressor'
    rows=[]
    ranking=destination/'post_ebm_ranking.csv'
    previous=pd.read_csv(ranking) if ranking.is_file() else pd.DataFrame(columns=['variant'])
    for family,variant in SCORING_EXPORTS:
        run=args.output/'1_Intermediate_Calculations'/f'{algorithm}__{family}'
        if not verify_run(run):
            # Keep the earlier ranking of a variant whose source model is absent, matching its saved export.
            rows.extend(previous[previous.variant.eq(variant)].to_dict('records'))
            continue
        destination.mkdir(parents=True,exist_ok=True)
        bundle=joblib.load(run/'model.joblib');names=bundle['features']
        importances=np.asarray(bundle['model'].term_importances(),dtype=float)
        for rank,index in enumerate(np.argsort(-importances),1):
            spec=definitions[names[index]]
            rows.append({'variant':variant,'source':spec['source'],'feature':spec['original_feature'],
                         'mean_normalized_importance':float(importances[index]/importances.sum()) if importances.sum() else 0.,
                         'selected_for_final_model':'yes','final_rank':rank,'selection_basis':'Training-eligible features; not validation-selected'})
        descriptors=[{'source':definitions[n]['source'],'feature':definitions[n]['original_feature']} for n in names]
        signature={'dataset_id':token(manifest),'source_completed_sha256':digest(run/'completed.json')}
        stamp=destination/f'export_{variant}.json'
        exported=destination/f'ebm_model_{variant}.joblib'
        try:
            saved=read_json(stamp) if stamp.is_file() else None
        except json.JSONDecodeError:
            saved=None
        expected={**signature,'model_sha256':digest(exported)} if exported.is_file() else None
        if expected is None or saved!=expected:
            descriptor,temporary=tempfile.mkstemp(prefix='.ebm_export_',suffix='.joblib',dir=destination)
            os.close(descriptor)
            temporary=Path(temporary)
            try:
                joblib.dump({'model':bundle['model'],'features':descriptors,'dataset_id':token(manifest),
                             'task':args.task,'algorithm':algorithm},temporary,compress=3)
                signature['model_sha256']=digest(temporary)
                os.replace(temporary,exported)
                write_json(stamp,signature)
            finally:
                temporary.unlink(missing_ok=True)
    if rows:
        content=pd.DataFrame(rows).to_csv(index=False)
        if not ranking.is_file() or ranking.read_text(encoding='utf-8',newline='')!=content:
            with open(ranking,'w',encoding='utf-8',newline='') as handle:handle.write(content)
    remove_old_report_files(destination,[destination/f'{kind}_{variant}.{extension}' for variant in FORMER_SCORING_EXPORTS
                                         for kind,extension in (('ebm_model','joblib'),('export','json'))])


def paired_rank_one_comparison(table, resamples_of, metric='macro_f1'):
    """Each model's difference from the rank-1 model of its leaderboard, on identical learner resamples."""
    columns = [f'{metric}_difference_from_rank_1', f'{metric}_difference_standard_deviation', f'{metric}_difference_95_low',
               f'{metric}_difference_95_high', 'distinguishable_from_rank_1']
    result = pd.DataFrame(index=table.index, columns=columns, dtype=object)
    for _, group in table.groupby(['evaluated_target', 'split'], sort=False):
        best = group.loc[group.quality_rank.idxmin()]
        reference = resamples_of(best)[best.evaluated_target][metric]
        for index, row in group.iterrows():
            differences = resamples_of(row)[row.evaluated_target][metric] - reference
            low, high = np.percentile(differences[np.isfinite(differences)], [2.5, 97.5])
            verdict = 'rank 1' if index == best.name else 'yes' if low > 0 or high < 0 else 'no'
            result.loc[index] = [row[metric] - best[metric], spread(differences), low, high, verdict]
    for column in columns[:-1]:
        result[column] = pd.to_numeric(result[column])
    return result


def summarize_baselines(args):
    """Rank completed models separately by evaluated target, preserving split labels."""
    manifest, frame = verify_shared(args.shared)
    comparisons, configurations = [], {}
    for task, folder in [('cefr', args.cefr_output), ('level', args.level_output)]:
        rows = collect_reports(folder, manifest, task, frame)[0]
        comparisons.extend(rows)
        for row in rows:
            path = folder/'1_Intermediate_Calculations'/row['experiment']
            configurations[(task, row['experiment'])] = path
    if not comparisons:raise ValueError('No completed models are available for the baseline summary.')

    def resamples_of(row):
        return run_resamples(configurations[(row['task'], row['experiment'])], row['split'], frame, manifest, row['task'])
    table = pd.DataFrame(comparisons).sort_values(
        ['evaluated_target','split','macro_f1','accuracy','ordinal_mae','task','experiment'],
        ascending=[True,True,False,False,True,True,True])
    table.insert(0,'quality_rank',table.groupby(['evaluated_target','split']).cumcount()+1)
    table['model_file'] = [manifest_path(configurations[(row.task,row.experiment)]/'model.joblib')
                           for row in table.itertuples()]
    ranked = paired_rank_one_comparison(table, resamples_of)
    for offset, column in enumerate(ranked.columns, table.columns.get_loc('macro_f1_standard_deviation') + 1):
        table.insert(offset, column, ranked[column])
    validation = table[table.split.eq('validation')].drop(columns=ranked.columns)
    baseline = validation[(validation.algorithm.eq('classifier') & validation.feature_set.eq('majority')) |
                          (validation.algorithm.eq('regressor') & validation.feature_set.eq('median'))]
    keys = ['task','algorithm','evaluated_target']
    baseline = baseline[keys+['experiment','macro_f1','accuracy','ordinal_mae']].rename(
        columns={'experiment':'baseline_experiment'})
    baseline['baseline_model'] = np.where(baseline.algorithm.eq('classifier'), 'Majority classifier', 'Median regressor')
    improvements = validation.merge(baseline, on=keys,how='left',validate='many_to_one',suffixes=('','_baseline'))
    for metric in ('macro_f1','accuracy'):
        improvements[metric+'_gain_over_baseline'] = improvements[metric]-improvements[metric+'_baseline']
    improvements['ordinal_mae_reduction_over_baseline'] = improvements.ordinal_mae_baseline-improvements.ordinal_mae
    # Paired: the model and its baseline are scored on the same resampled learners.
    for name, metric, sign in (('macro_f1_gain_over_baseline', 'macro_f1', 1), ('accuracy_gain_over_baseline', 'accuracy', 1),
                               ('ordinal_mae_reduction_over_baseline', 'ordinal_mae', -1)):
        spreads = []
        for _, row in improvements.iterrows():
            if pd.isna(row['baseline_experiment']):
                spreads.append(None)
                continue
            own = resamples_of(row)[row['evaluated_target']][metric]
            base = resamples_of({**row, 'experiment': row['baseline_experiment']})[row['evaluated_target']][metric]
            spreads.append(spread(sign * (own - base)))
        improvements.insert(improvements.columns.get_loc(name) + 1, name + '_standard_deviation', spreads)
    args.baseline_output.mkdir(parents=True,exist_ok=True)
    cefr = table[table.evaluated_target.eq('cefr')].copy()
    cefr = cefr.drop(columns=[name for name in cefr if name.startswith('regression_')])
    cefr['prediction_route'] = np.where(cefr.task.eq('cefr'), 'Direct CEFR classification',
        np.where(cefr.algorithm.eq('classifier'), 'Level probabilities summed into CEFR bands',
                 'Numerical level rounded and mapped to CEFR'))
    outputs={'1_Leaderboards/cefr_leaderboard.csv':cefr,
             '1_Leaderboards/level_leaderboard.csv':table[table.evaluated_target.eq('level')],
             '2_Model_Comparisons/baseline_improvements.csv':improvements}
    for name, frame in outputs.items():
        temporary=args.baseline_output/(name+'.tmp')
        temporary.parent.mkdir(parents=True,exist_ok=True)
        frame.to_csv(temporary,index=False)
        os.replace(temporary,args.baseline_output/name)
    # Remove superseded report products after the replacements have been saved.
    remove_old_report_files(args.baseline_output, [args.baseline_output/'2_Feature_Set_Comparisons'/name
        for name in ('baseline_improvements.csv','feature_selection_comparison.csv')] +
        [args.baseline_output/'2_Model_Comparisons/mi_vs_ig_model_comparison.csv'])
    for name in ('model_leaderboard.csv','best_models.csv','experiment_status.csv','summary_manifest.json',
                 'cefr_leaderboard.csv','level_leaderboard.csv','baseline_improvements.csv','feature_selection_comparison.csv'):
        obsolete = args.baseline_output/name
        if obsolete.is_file():
            if not obsolete.resolve().is_relative_to(args.baseline_output.resolve()):
                raise ValueError('Obsolete report path escapes the baseline output folder.')
            obsolete.unlink()
    remove_old_report_files(args.llm_output, [args.llm_output/'post_training_resources.csv'])
    log(f'Saved baseline comparisons: {args.baseline_output}')


def default_metadata():
    previous=BASE/'9_Previous_Runs/CEFR_Classification/0_Organizational_Files/splits.csv'
    return previous if previous.exists() else ROOT/'0_Data/Final database (main prompts).xlsx'


def main(task='cefr',argv=None):
    parser=argparse.ArgumentParser(description='Shared EBM experiments: fixed learners, CEFR classification, level classification and regression.')
    parser.add_argument('command',nargs='?',default='run',choices=['run','prepare','rank','train','report','summarize','final-test','test'])
    parser.add_argument('--shared',type=Path,default=SHARED)
    parser.add_argument('--rebuild-setup',action='store_true',help='Archive old setup and generated outputs before preparing rebuilt data.')
    parser.add_argument('--data',type=Path,default=DATA)
    parser.add_argument('--metadata',type=Path,default=default_metadata())
    parser.add_argument('--output',type=Path)
    parser.add_argument('--cefr-output',type=Path,default=BASE/'2_CEFR_Classification')
    parser.add_argument('--level-output',type=Path,default=BASE/'3_Level_Prediction')
    parser.add_argument('--ranking',type=Path,default=RANKING)
    parser.add_argument('--baseline-output',type=Path)
    parser.add_argument('--llm-output',type=Path)
    parser.add_argument('--seed',type=int,default=42)
    parser.add_argument('--sample-per-level',type=int,default=0)
    parser.add_argument('--rounds',type=int,default=200)
    parser.add_argument('--workers',type=int,default=2)
    parser.add_argument('--min-observations',type=int,default=20)
    parser.add_argument('--memory',default='auto')
    parser.add_argument('--experiments',nargs='+')
    parser.add_argument('--algorithms',nargs='+',choices=['classifier','regressor'])
    parser.add_argument('--use-ranked',action='store_true')
    parser.add_argument('--top-k',type=int,default=200)
    parser.add_argument('--runs',nargs='+')
    args=parser.parse_args(argv);args.task=task
    for attr in ('shared','data','metadata','ranking','cefr_output','level_output'):
        setattr(args,attr,getattr(args,attr).resolve())
    args.baseline_output=(args.baseline_output or args.shared.parent/'4_Baselines').resolve()
    args.llm_output=(args.llm_output or args.shared.parent/'5_Data_For_LLM_Post_Training').resolve()
    if min(args.rounds,args.workers,args.min_observations,args.top_k)<1 or args.sample_per_level<0:parser.error('Invalid nonpositive setting.')
    if task=='cefr' and args.algorithms and args.algorithms!=['classifier']:parser.error('CEFR task uses classifiers only.')
    args.output=(args.output or (args.cefr_output if task=='cefr' else args.level_output)).resolve()
    if task=='cefr':args.cefr_output=args.output
    else:args.level_output=args.output
    if args.command=='test':return run_tests()
    if args.command=='prepare':return prepare_shared(args)
    if args.command=='rank':
        if not (args.shared/'dataset_manifest.json').exists():prepare_shared(args)
        return feature_ranking(args)
    for name,folder in (('cefr',args.cefr_output),('level',args.level_output)):
        migrate_renamed_runs(folder,args.shared,name)
    if args.runs:
        renamed={old:new for renamed_task,old,new in RUN_RENAMES if renamed_task==args.task}
        args.runs=[renamed.get(name,name) for name in args.runs]
    if args.command=='train':return train_models(args)
    if args.command=='report':return write_reports(args)
    if args.command=='summarize':return summarize_baselines(args)
    if args.command=='final-test':return final_test(args)
    if not (args.shared/'dataset_manifest.json').exists():prepare_shared(args)
    if not (args.shared/'final_evaluation.json').exists():train_models(args)
    else:log('Final evaluation is sealed; refreshing saved reports only.')
    write_reports(args)

def run_tests():
    import shutil
    import unittest
    from types import SimpleNamespace
    from unittest.mock import patch
    module=sys.modules[__name__]
    cefr=module
    class WorkflowTests(unittest.TestCase):
        def setUp(self):
            # Small length lists for the synthetic features: the raw ERRANT counts are length-dependent,
            # and the total count stands in for the direct size measures of list 1.
            lists=Path(tempfile.mkdtemp());self.addCleanup(shutil.rmtree,lists,True)
            pd.DataFrame({'column':['accuracy__ERRANT__total_errors']}).to_csv(lists/'blacklist_1.csv',index=False)
            pd.DataFrame({'column':[f'accuracy__ERRANT__{kind}_errors' for kind in ('missing','unnecessary','replacement','other')]}).to_csv(
                lists/'blacklist_2.csv',index=False)
            patcher=patch.object(module,'LENGTH_BLACKLISTS',(lists/'blacklist_1.csv',lists/'blacklist_2.csv'))
            patcher.start();self.addCleanup(patcher.stop)
            # The missing-error count stands in for a duplicate feature.
            pd.DataFrame({'column':['accuracy__ERRANT__missing_errors']}).to_csv(lists/'blacklist_3.csv',index=False)
            patcher=patch.object(module,'DUPLICATE_LIST',lists/'blacklist_3.csv')
            patcher.start();self.addCleanup(patcher.stop)
        def fixture(self,root,polke=False):
            data=root/'data';data.mkdir();ids=np.arange(1,1801);level=(ids-1)%15+1
            cefr=np.array(LABELS)[(level-1)//3]
            table=pd.DataFrame({'text_id':ids,'cefr_level':cefr,'complexity__LCA__ld':level+ids/100000,
                                'accuracy__ERRANT__total_errors':10+ids%4,
                                'accuracy__ERRANT__missing_errors':ids%3,
                                'accuracy__ERRANT__unnecessary_errors':ids%2,
                                'accuracy__ERRANT__replacement_errors':5+ids%2,
                                'accuracy__ERRANT__other_errors':ids%2,
                                'accuracy__ERRANT__errors_per_100_words':(10+ids%4)/level})
            if polke:table.insert(3,'complexity__POLKE__polke_1_per_100_words',2.*level+ids%5/10)
            metadata=pd.DataFrame({'text_id':ids,'cefr_level':cefr,'level':level})
            with duckdb.connect() as con:
                con.register('features',table);con.register('metadata',metadata)
                con.execute(f'COPY features TO {literal(data / "feature_dataframe.parquet")} (FORMAT PARQUET)')
                con.execute(f'COPY metadata TO {literal(data / "text_metadata.parquet")} (FORMAT PARQUET)')
            kinds={'_errors':'error_count','_per_100_words':'error_rate_per_100_words'}
            specs=[{'column':column,'group':column.split('__')[0],'source':column.split('__')[1],'original_feature':column.split('__')[2],
                    'measure_type':next((kind for end,kind in kinds.items() if column.startswith('accuracy__') and column.endswith(end)),'test')}
                   for column in table.columns[2:]]
            pd.DataFrame(specs).to_csv(data/'feature_dictionary.csv',index=False)
            learners=metadata[['text_id','cefr_level']].copy();learners['learner_id']=((ids-1)//2).astype(str)
            learners.to_csv(root/'learners.csv',index=False)
            return ['--shared',str(root/'shared'),'--data',str(data),'--metadata',str(root/'learners.csv'),
                    '--cefr-output',str(root/'cefr'),'--level-output',str(root/'level'),'--ranking',str(root/'ranking'),
                    '--rounds','2','--workers','1','--min-observations','2']
        def test_complete_workflow(self):
            with tempfile.TemporaryDirectory() as temporary:
                root=Path(temporary);args=self.fixture(root)
                main('cefr',['prepare',*args]);manifest,splits=verify_shared(root/'shared')
                self.assertNotIn('complexity_without_text_size',read_json(root/'shared/feature_experiments.json'))
                self.assertEqual(splits.groupby('learner_id').split.nunique().max(),1)
                self.assertEqual(set(splits.level),set(range(1,16)))
                original_loader=load_matrix
                training_ids=set(splits.loc[splits.split=='train','text_id'])
                def check_training(*pos,**kw):
                    self.assertTrue(set(pos[1].text_id).issubset(training_ids))
                    return original_loader(*pos,**kw)
                with patch.object(module,'load_matrix',side_effect=check_training):main('cefr',['rank',*args,'--top-k','3'])
                main('cefr',['train',*args,'--use-ranked'])
                main('level',['train',*args,'--use-ranked'])
                main('cefr',['train',*args,'--use-ranked'])
                for task in ('cefr','level'):
                    legacy = root/task/'2_Prediction_Performance/0_All_Experiments/feature_set_comparison.csv'
                    legacy.parent.mkdir(parents=True,exist_ok=True)
                    legacy.write_text('old report\n',encoding='utf-8')
                    flat = [root/task/section/f'04_classifier__combined_without_length/{name}' for section,name in
                            (('2_Prediction_Performance',f'performance_metrics_{task}_04.csv'),
                             ('3_Feature_Contributions',f'Per_{"CEFR" if task=="cefr" else "Level"}/feature_contributions_{task}_04_1.csv'))]
                    for path in flat:
                        path.parent.mkdir(parents=True,exist_ok=True);path.write_text('earlier flat layout\n',encoding='utf-8')
                    exports = root/'5_Data_For_LLM_Post_Training'/('1_CEFR_Reference_Profiles' if task=='cefr' else '2_Level_Reference_Profiles')/'3_Reusable_Models_For_LLMs'
                    exports.mkdir(parents=True,exist_ok=True)
                    former = [exports/f'{kind}_{variant}.{extension}' for variant in FORMER_SCORING_EXPORTS
                              for kind,extension in (('ebm_model','joblib'),('export','json'))]
                    for path in former:path.write_text('old export\n',encoding='utf-8')
                    main(task,['report',*args])
                    self.assertFalse(legacy.exists())
                    self.assertFalse(any(path.exists() for path in former))
                    self.assertTrue((exports/'ebm_model_without_length_correlated_features.joblib').is_file())
                    self.assertEqual(set(pd.read_csv(exports/'post_ebm_ranking.csv').variant),{'all_features','without_length_correlated_features'})
                for task in ('cefr','level'):
                    performance = root/task/'2_Prediction_Performance'
                    comparison = pd.read_csv(performance/f'0_All_Experiments/performance_metrics_{task}_all_experiments.csv')
                    primary = comparison[comparison.evaluated_target.eq(task) & comparison.split.eq('validation')]
                    quality = primary.set_index('experiment').quality_rank.to_dict()
                    folders = ranked_experiment_folders(quality,primary.experiment,task)
                    # Grouped folders; the level task first splits classifiers and regressors.
                    tops = {'cefr':{'0_All_Experiments','1_Top_Performing','2_Length_Features_Filtered',
                                    '3_Mutual_Information_And_Information_Gain','5_Simple_Models','6_Duplicate_Features_Filtered'},
                            'level':{'0_All_Experiments','1_Classifier','2_Regressor'}}[task]
                    for section in (performance,root/task/'3_Feature_Contributions'):
                        self.assertEqual({path.name for path in section.iterdir() if path.is_dir()},tops)
                    prefix = '1_Classifier/' if task=='level' else ''
                    self.assertEqual(folders['classifier__combined_without_length_features'],
                        f"{prefix}2_Length_Features_Filtered/1_Without_Length_Features/{quality['classifier__combined_without_length_features']:02d}_combined")
                    self.assertEqual(folders['classifier__training_mutual_information_features'],
                        f"{prefix}3_Mutual_Information_And_Information_Gain/{quality['classifier__training_mutual_information_features']:02d}_mutual_information")
                    if task=='level':
                        self.assertEqual(folders['regressor__median'],f"2_Regressor/5_Simple_Models/{quality['regressor__median']:02d}_median")
                    self.assertTrue((root/task/'1_Intermediate_Calculations/classifier__combined/feature_columns.json').is_file())
                    contribution_root = root/task/'3_Feature_Contributions'
                    combined = pd.read_csv(contribution_root/f'0_All_Experiments/feature_contributions_{task}_all_experiments.csv', dtype={'output':str})
                    runs = list((root/task/'1_Intermediate_Calculations').glob('*/completed.json'))
                    for marker in runs:
                        name = marker.parent.name
                        separate = pd.read_csv(contribution_root/folders[name]/grouped_report_name('feature_contributions.csv',folders[name],task), dtype={'output':str})
                        pd.testing.assert_frame_equal(separate, combined[combined.experiment.eq(name)].reset_index(drop=True), check_dtype=False)
                        self.assertFalse((contribution_root/folders[name]/'post_model_feature_ranking.csv').exists())
                        self.assertIn('overall_feature_rank',separate)
                        separate_metrics = pd.read_csv(performance/folders[name]/grouped_report_name('performance_metrics.csv',folders[name],task))
                        pd.testing.assert_frame_equal(separate_metrics,comparison[comparison.experiment.eq(name)].reset_index(drop=True),check_dtype=False)
                        # One feature table per CEFR band or course level, for classifiers with features only.
                        source = marker.parent/'feature_contributions.csv'  # baselines have no features
                        saved = pd.read_csv(source, dtype={'output':str}) if source.is_file() else pd.DataFrame(columns=['feature','output'])
                        subfolder = contribution_root/folders[name]/('Per_CEFR' if task=='cefr' else 'Per_Level')
                        number = Path(folders[name]).name.split('_',1)[0]
                        outputs = [] if name.startswith('regressor__') or saved.empty else sorted(saved.output.unique())
                        tables = sorted(path.name for path in subfolder.iterdir()) if subfolder.is_dir() else []
                        self.assertEqual(tables, sorted(f'feature_contributions_{task}_{number}_{output}.csv' for output in outputs))
                        self.assertEqual(sorted(path.name for path in (contribution_root/folders[name]).iterdir()),
                                         sorted([grouped_report_name('feature_contributions.csv',folders[name],task),
                                                 grouped_report_name('tool_contributions.csv',folders[name],task)]+([subfolder.name] if outputs else [])))
                        for output in outputs:
                            table = pd.read_csv(subfolder/f'feature_contributions_{task}_{number}_{output}.csv')
                            self.assertEqual(list(table.columns), ['feature_name','contribution'])
                            expected = saved[saved.output.eq(output)].sort_values(['mean_absolute_contribution','feature'], ascending=[False,True])
                            self.assertEqual(list(table.feature_name), list(expected.feature))
                            self.assertTrue(np.allclose(table.contribution, expected.mean_absolute_contribution))
                    labels = {'cefr': manifest['labels'], 'level': [str(level) for level in range(1,16)]}[task]
                    combined_tables = [path for path in (contribution_root/folders['classifier__combined']).rglob('*.csv')
                                       if OUTPUT_TABLE_NAME.fullmatch(path.name)]
                    self.assertEqual(sorted(path.stem.rsplit('_',1)[1] for path in combined_tables), sorted(labels))
                    self.assertEqual({path.parent.name for path in combined_tables}, {'Per_CEFR' if task=='cefr' else 'Per_Level'})
                    self.assertFalse((contribution_root/'Per_CEFR').exists() or (contribution_root/'Per_Level').exists())
                    # Tool tables: shares add up to one per output, and a one-feature tool equals that feature.
                    tools=pd.read_csv(contribution_root/f'0_All_Experiments/tool_contributions_{task}_all_experiments.csv',dtype={'output':str})
                    for name in folders:
                        rows=tools[tools.experiment.eq(name)]
                        columns=read_json(root/task/'1_Intermediate_Calculations'/name/'feature_columns.json')
                        if not columns:
                            self.assertTrue(rows.empty);continue
                        self.assertTrue(np.allclose(rows.groupby('output').share.sum(),1))
                        self.assertEqual(rows.output.iloc[0],'overall')
                        self.assertEqual(dict(rows.drop_duplicates('tool').set_index('tool').features),dict(Counter(c.split('__')[1] for c in columns)))
                        self.assertTrue((root/task/'1_Intermediate_Calculations'/name/'tool_contributions.json').is_file())
                        if [c for c in columns if '__LCA__' in c]==['complexity__LCA__ld']:
                            saved=pd.read_csv(root/task/'1_Intermediate_Calculations'/name/'feature_contributions.csv',dtype={'output':str})
                            lca=rows[rows.tool.eq('LCA') & rows.output.ne('overall')].set_index('output').contribution.sort_index()
                            feature=saved[saved.feature.eq('complexity__LCA__ld')].set_index('output').mean_absolute_contribution.sort_index()
                            self.assertTrue(np.allclose(lca,feature))
                with patch.object(module,'load_matrix',side_effect=AssertionError('Must reuse the cached tool contributions')):
                    main('level',['report',*args])
                    self.assertFalse((root/task/'2_Prediction_Performance/2_Feature_Contributions').exists())
                    for path in performance.rglob('per_*_metrics_*.csv'):
                        table=pd.read_csv(path)
                        if 'precision' in table:
                            self.assertIn('accuracy',table)
                            present=table.support.gt(0)
                            self.assertTrue(np.allclose(table.loc[present,'accuracy'],table.loc[present,'recall']))
                            self.assertTrue(table.loc[~present,'accuracy'].isna().all())
                # Excel cannot open two files with the same name at once, so every report name must be unique.
                names=[path.name for task in ('cefr','level') for section in ('2_Prediction_Performance','3_Feature_Contributions')
                       for path in (root/task/section).rglob('*.csv')]
                self.assertEqual(len(names),len(set(names)))
                self.assertIn('performance_metrics_cefr_01.csv',names);self.assertIn('performance_metrics_level_all_experiments.csv',names)
                comp=pd.read_csv(root/'cefr/1_Intermediate_Calculations/cefr_prediction_routes_comparison.csv')
                self.assertEqual(len(comp),45)
                self.assertFalse(list(root.glob('*/1_Intermediate_Calculations/*complexity_without_text_size')))
                self.assertEqual(len(list(root.glob('*/1_Intermediate_Calculations/*_without_length_features'))),3+3)
                self.assertEqual(len(list(root.glob('*/1_Intermediate_Calculations/*_without_length_correlated_features'))),3+3)
                prediction=pd.read_csv(root/'level/1_Intermediate_Calculations/classifier__combined/validation_predictions.csv')
                self.assertEqual(set(prediction.text_id),set(splits.loc[splits.split=='validation','text_id']))
                for label in manifest['labels']:
                    cols=[f'probability_level_{level}' for level,cefr in manifest['level_to_cefr'].items() if cefr==label]
                    self.assertTrue(np.allclose(prediction[cols].sum(axis=1),prediction[f'probability_cefr_{label}']))
                # Exercise the relocated profile scorer with these same synthetic learners.
                import importlib.util
                profile_path=BASE/'5_Data_For_LLM_Post_Training/2_Level_Reference_Profiles/build_level_profiles.py'
                spec=importlib.util.spec_from_file_location('profile_workflow_test',profile_path)
                profile=importlib.util.module_from_spec(spec);spec.loader.exec_module(profile)
                (root/'2_Explainable_Boosting_Machine').mkdir()
                raw=splits[['writing_id','level','cefr_level','learner_id']].rename(columns={'cefr_level':'cefr'})
                raw['ld']=raw.level + raw.writing_id/100000
                raw.to_csv(root/'lca.csv',index=False)
                with patch.object(profile,'dataframe_inputs',return_value={'LCA':root/'lca.csv'}):
                    profile.main(['--thesis-root',str(root),'--ebm-dir',str(root/'5_Data_For_LLM_Post_Training/2_Level_Reference_Profiles/3_Reusable_Models_For_LLMs'),
                                  '--output-dir',str(root/'5_Data_For_LLM_Post_Training/2_Level_Reference_Profiles'),'--split-file',str(root/'shared/splits.csv'),
                                  '--min-observations','2','--min-level-texts','2','--min-profile-observations','2'])
                self.assertTrue((root/'5_Data_For_LLM_Post_Training/2_Level_Reference_Profiles/3_Reusable_Models_For_LLMs/weights_ebm_selected_features_without_length_correlated.json').exists())
                self.assertEqual(len(list((root/'5_Data_For_LLM_Post_Training/2_Level_Reference_Profiles/3_Reusable_Models_For_LLMs').glob('weights_*.json'))),4)
                self.assertTrue((root/'5_Data_For_LLM_Post_Training/2_Level_Reference_Profiles/2_Profile_Validation/feature_review.csv').is_file())
                exported=joblib.load(root/'5_Data_For_LLM_Post_Training/2_Level_Reference_Profiles/3_Reusable_Models_For_LLMs/ebm_model_without_length_correlated_features.joblib')
                self.assertEqual(exported['features'],[{'source':'LCA','feature':'ld'}])
                self.assertTrue(np.isfinite(exported['model'].predict([[5.]])).all())
                with patch.object(profile,'dataframe_inputs',return_value={'LCA':root/'lca.csv'}):
                    profile.main(['--profile-target','cefr','--thesis-root',str(root),
                        '--ebm-dir',str(root/'5_Data_For_LLM_Post_Training/1_CEFR_Reference_Profiles/3_Reusable_Models_For_LLMs'),
                        '--output-dir',str(root/'5_Data_For_LLM_Post_Training/1_CEFR_Reference_Profiles'),'--split-file',str(root/'shared/splits.csv'),
                        '--min-observations','2','--min-level-texts','2','--min-profile-observations','2'])
                cefr_weights=profile.load_profile_weights(root/'5_Data_For_LLM_Post_Training/1_CEFR_Reference_Profiles/3_Reusable_Models_For_LLMs/weights_all_features.json')
                self.assertEqual(set(cefr_weights['distance_methods']),{'cefr'})
                self.assertEqual({item['target'] for item in cefr_weights['profiles']['cefr']},set(manifest['labels']))
                self.assertFalse(cefr_weights['profiles']['level'])
                self.assertFalse(list((root/'5_Data_For_LLM_Post_Training/1_CEFR_Reference_Profiles').rglob('*.xlsx')))
                self.assertTrue(comp.groupby('split').quality_rank.apply(lambda x:list(x)==list(range(1,len(x)+1))).all())
                profile_splits=pd.read_csv(root/'5_Data_For_LLM_Post_Training/2_Level_Reference_Profiles/1_Intermediate_Calculations/1_Data_Splits/split_assignments.csv')
                self.assertEqual(set(profile_splits.loc[profile_splits.split=='validation','writing_id']),set(splits.loc[splits.split=='validation','text_id']))
                (root/'4_Baselines').mkdir(exist_ok=True)
                for obsolete in ('best_models.csv','model_leaderboard.csv','experiment_status.csv','summary_manifest.json',
                                 'cefr_leaderboard.csv','level_leaderboard.csv','baseline_improvements.csv','feature_selection_comparison.csv'):
                    (root/'4_Baselines'/obsolete).write_text('previous generated output')
                removed_comparison=root/'4_Baselines/2_Model_Comparisons/mi_vs_ig_model_comparison.csv'
                removed_comparison.parent.mkdir(parents=True,exist_ok=True)
                removed_comparison.write_text('previous generated comparison')
                removed_index=root/'5_Data_For_LLM_Post_Training/post_training_resources.csv'
                removed_index.write_text('previous generated index')
                main('cefr',['summarize',*args])
                for target,count in [('cefr',45),('level',30)]:
                    board=pd.read_csv(root/f'4_Baselines/1_Leaderboards/{target}_leaderboard.csv')
                    self.assertEqual(set(board.evaluated_target),{target})
                    self.assertEqual(len(board),count)
                    self.assertEqual(list(board.quality_rank),list(range(1,count+1)))
                    self.assertTrue(board.macro_f1.is_monotonic_decreasing)
                    self.assertEqual(list(board.columns[board.columns.get_loc('macro_f1')+1:][:6]),
                                     ['macro_f1_standard_deviation','macro_f1_difference_from_rank_1','macro_f1_difference_standard_deviation',
                                      'macro_f1_difference_95_low','macro_f1_difference_95_high','distinguishable_from_rank_1'])
                    self.assertTrue(board[[f'{name}_standard_deviation' for name in CLASS_METRICS]].notna().all().all())
                    self.assertEqual(list(board.distinguishable_from_rank_1),['rank 1']+['yes' if low>0 or high<0 else 'no'
                                     for low,high in zip(board.macro_f1_difference_95_low[1:],board.macro_f1_difference_95_high[1:])])
                    self.assertTrue(np.allclose(board.macro_f1_difference_from_rank_1,board.macro_f1-board.macro_f1.iloc[0]))
                    self.assertTrue((board.macro_f1_difference_95_low<=board.macro_f1_difference_95_high).all())
                level_board=pd.read_csv(root/'4_Baselines/1_Leaderboards/level_leaderboard.csv')
                regressors=level_board[level_board.algorithm.eq('regressor') & level_board.feature_set.ne('median')]
                self.assertTrue(regressors[[f'regression_{name}_standard_deviation' for name in REGRESSION_METRICS]].gt(0).all().all())
                self.assertTrue(level_board.loc[level_board.feature_set.eq('median'),'regression_spearman_correlation_standard_deviation'].isna().all())
                per_level=pd.read_csv(root/'level/2_Prediction_Performance/0_All_Experiments/per_level_metrics_level_all_experiments.csv')
                self.assertEqual(list(per_level.columns[per_level.columns.get_loc('accuracy'):][:9]),
                                 ['accuracy','accuracy_standard_deviation','precision','precision_standard_deviation','recall','recall_standard_deviation','f1-score','f1-score_standard_deviation','support'])
                course=pd.read_csv(root/'cefr/2_Prediction_Performance/0_All_Experiments/per_level_metrics_cefr_all_experiments.csv')
                self.assertTrue(course[[f'{name}_standard_deviation' for name in COURSE_LEVEL_METRICS]].notna().all().all())
                for name in ('best_models.csv','model_leaderboard.csv','experiment_status.csv','summary_manifest.json',
                             'cefr_leaderboard.csv','level_leaderboard.csv','baseline_improvements.csv','feature_selection_comparison.csv'):
                    self.assertFalse((root/'4_Baselines'/name).exists())
                self.assertFalse(removed_comparison.exists())
                self.assertFalse(removed_index.exists())
                self.assertEqual(len(list((root/'5_Data_For_LLM_Post_Training').rglob('weights_*.json'))),8)
                self.assertEqual(len(list((root/'5_Data_For_LLM_Post_Training').rglob('ebm_model_*.joblib'))),4)
                gains=pd.read_csv(root/'4_Baselines/2_Model_Comparisons/baseline_improvements.csv')
                self.assertTrue(np.allclose(gains.accuracy_gain_over_baseline,gains.accuracy-gains.accuracy_baseline))
                for name in ('macro_f1_gain_over_baseline','accuracy_gain_over_baseline','ordinal_mae_reduction_over_baseline'):
                    self.assertEqual(gains.columns[gains.columns.get_loc(name)+1],name+'_standard_deviation')
                    self.assertTrue(np.allclose(gains.loc[gains.experiment.eq(gains.baseline_experiment),name+'_standard_deviation'],0))
                    self.assertTrue(gains[name+'_standard_deviation'].ge(0).all())
                    self.assertTrue(gains.loc[gains.feature_set.eq('combined'),name+'_standard_deviation'].gt(0).all())
                self.assertFalse(any(name.startswith('macro_f1_difference') for name in gains))
                self.assertTrue(np.allclose(gains.ordinal_mae_reduction_over_baseline,gains.ordinal_mae_baseline-gains.ordinal_mae))
                self.assertEqual(set(gains.loc[gains.algorithm.eq('regressor'),'baseline_experiment']), {'regressor__median'})
                self.assertEqual(set(gains.loc[gains.algorithm.eq('regressor'),'baseline_model']), {'Median regressor'})
                self.assertEqual(set(gains.loc[gains.algorithm.eq('classifier'),'baseline_experiment']), {'classifier__majority'})
                self.assertFalse(any('over_majority' in name for name in gains))
                self.assertTrue((root/'level/1_Intermediate_Calculations/regressor__median/model.joblib').is_file())
                self.assertFalse((root/'level/1_Intermediate_Calculations/regressor__majority').exists())
                classifications=pd.read_csv(root/'cefr/2_Prediction_Performance/0_All_Experiments/CEFR_classifications_cefr_all_experiments.csv')
                totals=classifications.groupby(['experiment','split','evaluated_target','actual']).percentage.sum()
                self.assertTrue(np.allclose(totals,100))
                main('cefr',['final-test',*args,'--runs','classifier__combined'])
                with self.assertRaisesRegex(ValueError,'sealed'):main('level',['train',*args])
                main('level',['final-test',*args,'--runs','classifier__combined','regressor__combined'])
                main('level',['final-test',*args,'--runs','classifier__combined','regressor__combined'])
                main('level',['report',*args])
                with (root/'shared/splits.csv').open('a') as f:f.write('\n')
                with self.assertRaisesRegex(ValueError,'splits.csv changed'):verify_shared(root/'shared')
        def test_contribution_reports_preserve_models_and_migrate(self):
            with tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                legacy = root/'2_Prediction_Performance/2_Feature_Contributions'
                legacy.mkdir(parents=True)
                for name in ('feature_contributions.csv', 'post_model_feature_ranking.csv'):
                    (legacy/name).write_text('old generated summary')
                effects = [
                    {'experiment':'classifier__combined', 'feature':'x', 'output':'A1', 'mean_absolute_contribution':1.},
                    {'experiment':'classifier__combined', 'feature':'x', 'output':'A2', 'mean_absolute_contribution':3.},
                    {'experiment':'classifier__combined', 'feature':'y', 'output':'A1', 'mean_absolute_contribution':4.},
                    {'experiment':'classifier__combined', 'feature':'y', 'output':'A2', 'mean_absolute_contribution':6.},
                    {'experiment':'regressor__combined', 'feature':'x', 'output':'numerical_level', 'mean_absolute_contribution':10.}]
                quality = {'classifier__combined':1, 'regressor__combined':2, 'classifier__majority':3}
                old_folder = root/'3_Feature_Contributions/01_classifier__combined'
                old_folder.mkdir(parents=True)
                for name in ('feature_contributions.csv','post_model_feature_ranking.csv',
                             'post_model_feature_ranking_01_classifier__combined.csv','feature_contributions_01_classifier__combined.csv'):
                    (old_folder/name).write_text('old generated table')
                (old_folder/'notes.txt').write_text('keep my notes')
                flat = root/'3_Feature_Contributions/Per_CEFR'
                flat.mkdir()
                (flat/'feature_contributions_classifier__combined_A1.csv').write_text('previous per-task layout')
                write_contribution_reports(root, effects, quality, quality, 'cefr')
                destination = root/'3_Feature_Contributions'
                self.assertFalse(flat.exists())
                self.assertFalse((old_folder/'feature_contributions_01_classifier__combined.csv').exists())
                details = pd.read_csv(destination/'1_Top_Performing/01_combined/feature_contributions_cefr_01.csv')
                ranking = details.drop_duplicates('feature')
                self.assertEqual(list(ranking.feature), ['y', 'x'])
                self.assertEqual(list(ranking.overall_mean_absolute_contribution), [5., 2.])
                self.assertEqual(list(ranking.overall_feature_rank), [1, 2])
                self.assertEqual(list(details.mean_absolute_contribution), [4., 6., 1., 3.])
                for label, values in (('A1', [4., 1.]), ('A2', [6., 3.])):
                    per_output = pd.read_csv(destination/f'1_Top_Performing/01_combined/Per_CEFR/feature_contributions_cefr_01_{label}.csv')
                    self.assertEqual(list(per_output.columns), ['feature_name', 'contribution'])
                    self.assertEqual(list(per_output.feature_name), ['y', 'x'])
                    self.assertEqual(list(per_output.contribution), values)
                self.assertEqual(sorted(path.name for path in (destination/'1_Top_Performing/01_combined/Per_CEFR').iterdir()),
                                 ['feature_contributions_cefr_01_A1.csv', 'feature_contributions_cefr_01_A2.csv'])
                self.assertFalse(list(destination.glob('Per_*')) + list(destination.glob('*/0[23]_*/Per_*')))
                regressor = pd.read_csv(destination/'1_Top_Performing/02_combined/feature_contributions_cefr_02.csv')
                self.assertEqual(regressor.mean_absolute_contribution.iloc[0], 10.)
                baseline = pd.read_csv(destination/'5_Simple_Models/03_majority/feature_contributions_cefr_03.csv')
                self.assertTrue(baseline.empty)
                combined_path = destination/'0_All_Experiments/feature_contributions_cefr_all_experiments.csv'
                self.assertEqual(len(pd.read_csv(combined_path)), len(effects))
                before = fingerprint(combined_path)
                write_contribution_reports(root, effects, quality, quality, 'cefr')
                self.assertEqual(before, fingerprint(combined_path))
                self.assertFalse(legacy.exists())
                self.assertFalse(list(destination.rglob('post_model_feature_ranking*.csv')))
                self.assertFalse((old_folder/'feature_contributions.csv').exists())
                self.assertTrue((old_folder/'notes.txt').is_file())
                (old_folder/'notes.txt').unlink()
                changed = {**quality,'classifier__combined':2,'regressor__combined':1}
                write_contribution_reports(root,effects,changed,changed,'cefr')
                self.assertTrue((destination/'1_Top_Performing/02_combined/feature_contributions_cefr_02.csv').is_file())
                self.assertEqual(sorted(path.name for path in (destination/'1_Top_Performing/02_combined/Per_CEFR').iterdir()),
                                 ['feature_contributions_cefr_02_A1.csv', 'feature_contributions_cefr_02_A2.csv'])
                self.assertFalse((destination/'01_classifier__combined').exists())
                write_contribution_reports(root, [row for row in effects if row['output'] != 'A2'], changed, changed, 'cefr')
                self.assertEqual([path.name for path in (destination/'1_Top_Performing/02_combined/Per_CEFR').iterdir()],
                                 ['feature_contributions_cefr_02_A1.csv'])
                self.assertTrue((destination/'1_Top_Performing/01_combined/feature_contributions_cefr_01.csv').is_file())
                self.assertFalse((destination/'1_Top_Performing/01_combined/Per_CEFR').exists())
                with self.assertRaisesRegex(ValueError, 'Invalid experiment'):
                    write_contribution_reports(root, [], {}, ['../outside'], 'cefr')
                with self.assertRaisesRegex(ValueError, 'escapes'):
                    save_report_tables(destination, {'../outside.csv': baseline})
                empty = root/'majority_only'
                write_contribution_reports(empty, [], {'classifier__majority':1}, ['classifier__majority'], 'cefr')
                self.assertTrue(pd.read_csv(empty/'3_Feature_Contributions/0_All_Experiments/feature_contributions_cefr_all_experiments.csv').empty)

        def test_accuracy_within_each_actual_label(self):
            # A1 precision is 2/3, its within-class accuracy is 2/4, and
            # one-versus-rest accuracy is 4/7. These must not be confused.
            metrics=class_metrics(['A1']*4+['A2']*3,['A1','A1','A2','B1','A1','A2','A2'],['A1','A2','B1'])
            rows=list(per_label_report_rows(metrics))
            self.assertEqual(rows[0]['accuracy'],.5)
            self.assertAlmostEqual(rows[0]['precision'],2/3)
            self.assertAlmostEqual(rows[1]['accuracy'],2/3)
            self.assertTrue(np.isnan(rows[2]['accuracy']))
            self.assertAlmostEqual(sum(row['accuracy']*row['support'] for row in rows if row['support'])/7,
                                   metrics['accuracy'])

        def test_aggregation_and_regression(self):
            rows=pd.DataFrame({'text_id':[1],'learner_id':['1'],'level':[2],'cefr_level':['A1']})
            manifest={'labels':['A1','A2','B1'],'level_to_cefr':{str(i):LABELS[(i-1)//3] for i in range(1,16)}}
            classifier=SimpleNamespace(classes_=np.array([1,2,4]),predict_proba=lambda x:np.array([[.3,.3,.4]]))
            metrics,pred=evaluate_model(classifier,np.zeros((1,1)),rows,manifest,'level','classifier')
            self.assertEqual(pred.predicted_level.iloc[0],4)
            self.assertEqual(pred.predicted_cefr.iloc[0],'A1')
            regressor=SimpleNamespace(predict=lambda x:np.array([2.2]))
            metrics,pred=evaluate_model(regressor,np.zeros((1,1)),rows,manifest,'level','regressor')
            self.assertEqual(pred.predicted_level.iloc[0],2)
            self.assertAlmostEqual(metrics['regression']['mae'],.2)
        def test_feature_and_split_guards(self):
            with self.assertRaisesRegex(ValueError,'Unknown feature'):feature_sql('cefr_level',set())
            frame=pd.DataFrame({'text_id':[1,2],'learner_id':['same','same'],'split':['train','validation'],'cefr_level':['A1','A1']})
            with self.assertRaises(ValueError):validate_splits(frame,['A1'])
        def test_median_baseline_migration_preserves_training_and_seal(self):
            for sealed in (False, True):
                with self.subTest(sealed=sealed), tempfile.TemporaryDirectory() as temporary:
                    root=Path(temporary);args=self.fixture(root)
                    main('cefr',['prepare',*args])
                    main('level',['train',*args,'--experiments','median','--algorithms','regressor'])
                    parent=root/'level/1_Intermediate_Calculations'
                    new=parent/'regressor__median';old=parent/'regressor__majority'
                    config=read_json(new/'run_config.json')
                    config.update(feature_set='majority', code_sha256=next(iter(COMPATIBLE_CODE_HASHES)),
                                  definition={'columns': [], 'rationale': 'Training majority baseline.'})
                    write_json(new/'run_config.json',config)
                    completed=read_json(new/'completed.json')
                    completed['configuration_id']=token(config)
                    completed['files']['run_config.json']=digest(new/'run_config.json')
                    write_json(new/'completed.json',completed)
                    self.assertTrue(all(path.resolve().is_relative_to(root.resolve()) for path in (new,old)))
                    new.rename(old)
                    before={p.name:digest(p) for p in old.iterdir()}
                    seal_path=root/'shared/final_evaluation.json'
                    if sealed:
                        write_json(seal_path,{'dataset_id':config['dataset_id'],
                            'locations':{'level':str((root/'level').resolve())},
                            'runs':{'level':{old.name:digest(old/'completed.json')}}})
                    migrate_renamed_runs(root/'level',root/'shared','level')
                    migrate_renamed_runs(root/'level',root/'shared','level')
                    self.assertFalse(old.exists())
                    self.assertEqual(before,{p.name:digest(p) for p in new.iterdir()})
                    desired=canonical_training_configuration(config)
                    desired['code_sha256']=digest(Path(__file__))
                    self.assertTrue(verify_run(new,desired))
                    with self.assertRaisesRegex(ValueError,'settings or training code changed'):
                        verify_run(new,{**desired,'rounds':config['rounds']+1})
                    if sealed:
                        self.assertEqual(read_json(seal_path)['runs']['level'],{new.name:before['completed.json']})
                        # A retry also completes a rename interrupted before writing the seal.
                        interrupted=read_json(seal_path)
                        interrupted['runs']['level']={old.name:before['completed.json']}
                        write_json(seal_path,interrupted)
                        migrate_renamed_runs(root/'level',root/'shared','level')
                        self.assertEqual(read_json(seal_path)['runs']['level'],{new.name:before['completed.json']})
                    else:
                        with patch.object(DummyRegressor,'fit',side_effect=AssertionError('Must reuse the trained median baseline')):
                            main('level',['train',*args,'--experiments','median','--algorithms','regressor'])
                    old.mkdir()
                    with self.assertRaisesRegex(ValueError,'Both regressor__majority and regressor__median exist'):
                        migrate_renamed_runs(root/'level',root/'shared','level')
                    old.rmdir()
                    (new/'model.joblib').write_bytes(b'changed artifact')
                    with self.assertRaisesRegex(ValueError,'Changed run artifact'):
                        migrate_renamed_runs(root/'level',root/'shared','level')

        def test_mutual_information_rename_reuses_trained_models(self):
            with tempfile.TemporaryDirectory() as temporary:
                root=Path(temporary);args=self.fixture(root)
                main('cefr',['prepare',*args]);main('cefr',['rank',*args,'--top-k','3'])
                for task in ('cefr','level'):
                    main(task,['train',*args,'--use-ranked','--experiments','training_mutual_information_features'])
                for task in ('cefr','level'):
                    parent=root/task/'1_Intermediate_Calculations'
                    for new in list(parent.iterdir()):
                        # Recreate a run saved under the former name by an earlier reviewed code version.
                        config=read_json(new/'run_config.json')
                        config.update(feature_set='training_ranked_features',code_sha256=next(iter(COMPATIBLE_CODE_HASHES)),
                                      definition={**config['definition'],'rationale':'Optional training-only ranked feature subset.'})
                        write_json(new/'run_config.json',config)
                        completed=read_json(new/'completed.json')
                        completed['configuration_id']=token(config)
                        completed['files']['run_config.json']=digest(new/'run_config.json')
                        write_json(new/'completed.json',completed)
                        new.rename(parent/new.name.replace('training_mutual_information_features','training_ranked_features'))
                snapshot=lambda:{(task,run.name):{path.name:digest(path) for path in run.iterdir()}
                                 for task in ('cefr','level') for run in (root/task/'1_Intermediate_Calculations').iterdir()}
                before=snapshot()
                self.assertEqual({name for _,name in before},{'classifier__training_ranked_features','regressor__training_ranked_features'})
                with patch.object(ExplainableBoostingClassifier,'fit',side_effect=AssertionError('Must reuse the renamed models')), \
                     patch.object(ExplainableBoostingRegressor,'fit',side_effect=AssertionError('Must reuse the renamed models')):
                    for task in ('cefr','level'):
                        main(task,['train',*args,'--use-ranked','--experiments','training_mutual_information_features'])
                self.assertEqual(snapshot(),{(task,name.replace('training_ranked_features','training_mutual_information_features')):files
                                             for (task,name),files in before.items()})
                main('level',['report',*args])
                performance=pd.read_csv(root/'level/2_Prediction_Performance/0_All_Experiments/performance_metrics_level_all_experiments.csv')
                self.assertEqual(set(performance.feature_set),{'training_mutual_information_features'})

        def test_length_rename_reuses_trained_models(self):
            with tempfile.TemporaryDirectory() as temporary:
                root=Path(temporary);args=self.fixture(root)
                main('cefr',['prepare',*args])
                experiments=['--experiments','complexity_without_length_correlated_features','combined_without_length_correlated_features']
                for task in ('cefr','level'):main(task,['train',*args,*experiments])
                for task in ('cefr','level'):
                    parent=root/task/'1_Intermediate_Calculations'
                    for new in list(parent.iterdir()):
                        # Recreate a run saved under the former name by the reviewed code version that trained it.
                        config=read_json(new/'run_config.json')
                        config.update(feature_set=config['feature_set'].replace('_correlated_features',''),
                                      code_sha256='54d298cbd052a147e321b283328211b68d23c9a0c73a52213c098b304b3e2048')
                        write_json(new/'run_config.json',config)
                        completed=read_json(new/'completed.json')
                        completed['configuration_id']=token(config)
                        completed['files']['run_config.json']=digest(new/'run_config.json')
                        write_json(new/'completed.json',completed)
                        new.rename(parent/new.name.replace('_correlated_features',''))
                snapshot=lambda:{(task,run.name):{path.name:digest(path) for path in run.iterdir()}
                                 for task in ('cefr','level') for run in (root/task/'1_Intermediate_Calculations').iterdir()}
                before=snapshot()
                self.assertEqual(len(before),2+4)
                with patch.object(ExplainableBoostingClassifier,'fit',side_effect=AssertionError('Must reuse the renamed models')), \
                     patch.object(ExplainableBoostingRegressor,'fit',side_effect=AssertionError('Must reuse the renamed models')):
                    for task in ('cefr','level'):main(task,['train',*args,*experiments])
                self.assertEqual(snapshot(),{(task,name+'_correlated_features'):files for (task,name),files in before.items()})

        def test_training_lock_waits_for_a_parallel_process(self):
            import threading
            with tempfile.TemporaryDirectory() as temporary:
                run=Path(temporary)/'classifier__combined';lock=run.parent/'.classifier__combined.lock'
                # A live process owns the run: wait until it releases the lock.
                lock.write_text(str(os.getppid()),encoding='ascii')
                release=threading.Timer(0.3,lambda:lock.unlink())
                release.start()
                started=time.time()
                with training_lock(run,poll_seconds=0.05):
                    self.assertGreaterEqual(time.time()-started,0.25)
                    self.assertEqual(lock.read_text(encoding='ascii'),str(os.getpid()))
                self.assertFalse(lock.exists())
                # Stale locks from this process or a stopped one are taken over at once.
                stopped=max(psutil.pids())+100000
                for owner in (os.getpid(),stopped):
                    lock.write_text(str(owner),encoding='ascii')
                    started=time.time()
                    with training_lock(run,poll_seconds=5):
                        self.assertLess(time.time()-started,1)
                    self.assertFalse(lock.exists())
                # An error while training releases the lock.
                with self.assertRaises(RuntimeError):
                    with training_lock(run):raise RuntimeError('fit failed')
                self.assertFalse(lock.exists())

        def test_polke_only_experiment(self):
            with tempfile.TemporaryDirectory() as temporary:
                root=Path(temporary);args=self.fixture(root,polke=True)
                main('cefr',['prepare',*args]);manifest,_=verify_shared(root/'shared')
                column='complexity__POLKE__polke_1_per_100_words'
                self.assertNotIn('polke',read_json(root/'shared/feature_experiments.json'))
                # The fixed thesis setup still lists the retired experiment; it must not be trained.
                legacy=root/'legacy_shared';shutil.copytree(root/'shared',legacy)
                experiments=read_json(legacy/'feature_experiments.json')
                write_json(legacy/'feature_experiments.json',{**experiments,'complexity_without_text_size':experiments['complexity']})
                self.assertNotIn('complexity_without_text_size',experiment_definitions(SimpleNamespace(shared=legacy,task='cefr',use_ranked=False),manifest))
                for task in ('cefr','level'):
                    definitions=experiment_definitions(SimpleNamespace(shared=root/'shared',task=task,use_ranked=False),manifest)
                    self.assertEqual(definitions['polke'],{'columns':[column],'rationale':POLKE_RATIONALE})
                    # Without length features, only the list-1 feature goes; without length-correlated features, both lists go.
                    self.assertEqual(definitions['complexity_without_length_features']['columns'],['complexity__LCA__ld',column])
                    self.assertEqual(definitions['combined_without_length_features']['columns'],
                                     ['complexity__LCA__ld',column]+[f'accuracy__ERRANT__{kind}_errors' for kind in ('missing','unnecessary','replacement','other')]
                                     +['accuracy__ERRANT__errors_per_100_words'])
                    self.assertEqual(definitions['complexity_without_length_correlated_features']['columns'],['complexity__LCA__ld',column])
                    self.assertEqual(definitions['combined_without_length_correlated_features']['columns'],
                                     ['complexity__LCA__ld',column,'accuracy__ERRANT__errors_per_100_words'])
                    # Duplicate filtering removes the listed duplicate from each base set.
                    errors=[f'accuracy__ERRANT__{kind}_errors' for kind in ('total','unnecessary','replacement','other')]
                    rate='accuracy__ERRANT__errors_per_100_words'
                    self.assertEqual(definitions['combined_without_duplicate_features']['columns'],['complexity__LCA__ld',column,*errors,rate])
                    self.assertEqual(definitions['combined_without_length_and_duplicate_features']['columns'],['complexity__LCA__ld',column,*errors[1:],rate])
                    self.assertEqual(definitions['combined_without_length_correlated_and_duplicate_features']['columns'],['complexity__LCA__ld',column,rate])
                    self.assertEqual(definitions['length_only']['columns'],['accuracy__ERRANT__total_errors'])  # list 1 in these tests
                    self.assertIn(column,definitions['complexity']['columns'])
                    main(task,['train',*args,'--experiments','polke'])
                expected={'cefr':['classifier__polke'],'level':['classifier__polke','regressor__polke']}
                for task,names in expected.items():
                    parent=root/task/'1_Intermediate_Calculations'
                    self.assertEqual(sorted(path.name for path in parent.iterdir()),names)
                    for name in names:
                        self.assertEqual(read_json(parent/name/'feature_columns.json'),[column])
                        self.assertEqual(read_json(parent/name/'run_config.json')['feature_set'],'polke')
                with patch.object(ExplainableBoostingClassifier,'fit',side_effect=AssertionError('Must reuse the POLKE models')), \
                     patch.object(ExplainableBoostingRegressor,'fit',side_effect=AssertionError('Must reuse the POLKE models')):
                    for task in ('cefr','level'):main(task,['train',*args,'--experiments','polke'])

        def test_bootstrap_matches_expanded_resamples(self):
            rng=np.random.default_rng(3);size=90
            level=rng.integers(1,16,size);guess=np.clip(level+rng.integers(-2,3,size),1,15)
            predictions=pd.DataFrame({'text_id':np.arange(1,size+1),'learner_id':(np.arange(size)//3).astype(str),'level':level,
                                      'cefr_level':np.array(LABELS)[(level-1)//3],'predicted_level':guess,
                                      'predicted_level_continuous':np.clip(guess+rng.normal(0,.4,size),1.,15.),
                                      'predicted_cefr':np.array(LABELS)[(guess-1)//3]})
            predictions.loc[:4,'predicted_level_continuous']=15.  # ties, as clipping creates
            manifest={'labels':LABELS[:5]}
            weights=rng.multinomial(size,np.full(size,1/size),size=6).astype(float)
            level_metrics=resample_metrics(weights,prepare_resampling(predictions,manifest,'level'))
            cefr_metrics=resample_metrics(weights,prepare_resampling(predictions,manifest,'cefr'))
            for index,multiplicity in enumerate(weights.astype(int)):
                sample=predictions.loc[np.repeat(np.arange(size),multiplicity)]
                for target,truth,predicted,labels in (('cefr','cefr_level','predicted_cefr',LABELS[:5]),
                                                     ('level','level','predicted_level',list(range(1,16)))):
                    expected=class_metrics(sample[truth],sample[predicted],labels)
                    for name in CLASS_METRICS:
                        self.assertAlmostEqual(level_metrics[target][name][index],expected[name],places=10)
                    for column,row in enumerate(per_label_report_rows(expected)):
                        for name in LABEL_METRICS:
                            value=level_metrics[target]['labels'][name][index,column]
                            if np.isnan(row[name]):self.assertTrue(np.isnan(value))
                            else:self.assertAlmostEqual(value,row[name],places=10)
                for name in CLASS_METRICS:self.assertEqual(cefr_metrics['cefr'][name][index],level_metrics['cefr'][name][index])
                truth,estimate=sample.level,sample.predicted_level_continuous
                regression=level_metrics['regression']
                self.assertAlmostEqual(regression['mae'][index],mean_absolute_error(truth,estimate),places=10)
                self.assertAlmostEqual(regression['rmse'][index],np.sqrt(mean_squared_error(truth,estimate)),places=10)
                self.assertAlmostEqual(regression['r_squared'][index],r2_score(truth,estimate),places=10)
                self.assertAlmostEqual(regression['spearman_correlation'][index],spearmanr(truth,estimate).statistic,places=10)
                table=sample.assign(distance=(sample.cefr_level.map(LABELS.index)-sample.predicted_cefr.map(LABELS.index)).abs())
                for column,level_value in enumerate(np.unique(predictions.level)):
                    group=table[table.level.eq(level_value)]
                    for name,values in (('cefr_accuracy',group.distance.eq(0)),('cefr_ordinal_mae',group.distance),
                                        ('within_one_cefr_accuracy',group.distance.le(1))):
                        value=cefr_metrics['course_level'][name][index,column]
                        if group.empty:self.assertTrue(np.isnan(value))
                        else:self.assertAlmostEqual(value,float(values.mean()),places=10)
            frame=predictions.assign(split='validation')
            first=learner_draws(frame,'validation');second=learner_draws(frame.sample(frac=1,random_state=1),'validation')
            self.assertIs(first,second)
            text_ids,learners,draws=first
            self.assertEqual(draws.shape,(BOOTSTRAP_RESAMPLES,30))
            self.assertTrue((draws.sum(axis=1)==30).all())
            self.assertTrue(np.array_equal(pd.factorize(learners)[0],np.arange(size)//3))
            self.assertIsNone(spread(np.r_[np.full(10,np.nan),np.ones(90)]))
            self.assertAlmostEqual(spread(np.r_[np.nan,np.arange(99.)]),float(np.arange(99.).std(ddof=1)))

        def test_previous_training_reuse_checks_settings(self):
            with tempfile.TemporaryDirectory() as temporary:
                run=Path(temporary)
                previous={'code_sha256':next(iter(COMPATIBLE_CODE_HASHES)), 'rounds':200}
                write_json(run/'run_config.json',previous)
                write_json(run/'completed.json',{'configuration_id':token(previous),
                    'files':{'run_config.json':digest(run/'run_config.json')}})
                current={**previous,'code_sha256':digest(Path(__file__))}
                self.assertTrue(verify_run(run,current))
                with self.assertRaisesRegex(ValueError,'settings or training code changed'):
                    verify_run(run,{**current,'rounds':201})
        def test_legacy_rankings_migrate_without_recalculation(self):
            with tempfile.TemporaryDirectory() as temporary:
                root=Path(temporary);args=self.fixture(root)
                main('cefr',['prepare',*args]);manifest,splits=verify_shared(root/'shared')
                hashes={}
                for method,folder in [('mi',root/'ranking'),('ig',root/'ranking/2_Information_Gain')]:
                    folder.mkdir(parents=True,exist_ok=True)
                    column='mutual_information' if method=='mi' else 'information_gain_bits'
                    table=pd.DataFrame([{'target':target,'feature':'complexity__LCA__ld','eligible':True,
                        column:.5,'training_observations':100,'rank':1,'selected':True} for target in ('cefr','level')])
                    if method=='ig':table.insert(5,'best_split_threshold',.25)
                    table.to_csv(folder/'training_feature_ranking.csv',index=False)
                    config={'dataset_id':token(manifest),'method':RANKING_METHODS[method],
                        'code_sha256':next(iter(COMPATIBLE_CODE_HASHES)),'top_k':200,'min_observations':2,
                        'rows':int(splits.split.eq('train').sum())}
                    write_json(folder/'ranking_config.json',config)
                    write_json(folder/'completed.json',{'configuration_id':token(config),'files':{
                        name:digest(folder/name) for name in ('training_feature_ranking.csv','ranking_config.json')}})
                    hashes[method]=digest(folder/'training_feature_ranking.csv')
                with patch.object(module,'load_matrix',side_effect=AssertionError('Must reuse existing rankings')):
                    main('cefr',['rank',*args])
                    main('cefr',['rank',*args])
                config,tables=ranking_tables(root/'ranking',manifest)
                self.assertEqual(config['source_table_sha256'],hashes)
                self.assertEqual(set(tables),{'mi','ig'})
                self.assertTrue(all(table.selected.all() for table in tables.values()))
                self.assertFalse((root/'ranking/2_Information_Gain').exists())
                self.assertFalse((root/'ranking/training_feature_ranking.csv').exists())
                self.assertFalse((root/'ranking/ranking_config.json').exists())
                self.assertFalse((root/'ranking/completed.json').exists())
        def test_workbook_metadata_migration_and_integrity(self):
            import openpyxl
            with tempfile.TemporaryDirectory() as temporary:
                root=Path(temporary);args=self.fixture(root)
                main('cefr',['prepare',*args]);main('cefr',['rank',*args])
                manifest,_=verify_shared(root/'shared')
                folder=root/'ranking';path=folder/'training_feature_ranking.xlsx'
                config,tables=ranking_tables(folder,manifest)
                book=openpyxl.load_workbook(path)
                self.assertEqual(book['_Pipeline_Metadata'].sheet_state,'hidden')
                del book['_Pipeline_Metadata'];book.save(path);book.close()
                write_json(folder/'ranking_config.json',config)
                write_json(folder/'completed.json',{'configuration_id':token(config),'files':{
                    name:digest(folder/name) for name in ('training_feature_ranking.xlsx','ranking_config.json')}})
                with patch.object(module,'load_matrix',side_effect=AssertionError('Do not rerank saved features')):
                    main('cefr',['rank',*args]);main('cefr',['rank',*args])
                migrated,actual=ranking_tables(folder,manifest)
                self.assertEqual(config,migrated)
                for method in tables:pd.testing.assert_frame_equal(tables[method],actual[method])
                self.assertFalse(list(folder.glob('*.json')))
                book=openpyxl.load_workbook(path)
                book[RANKING_SHEET]['E2']=999
                book.save(path);book.close()
                with self.assertRaisesRegex(ValueError,'table changed'):ranking_tables(folder,manifest)
        def test_side_by_side_ranking_preserves_both_feature_selections(self):
            import openpyxl
            with tempfile.TemporaryDirectory() as temporary:
                root=Path(temporary);args=self.fixture(root)
                main('cefr',['prepare',*args])
                manifest,splits=verify_shared(root/'shared')
                folder=root/'ranking';folder.mkdir();path=folder/'training_feature_ranking.xlsx'
                features=['complexity__LCA__ld','accuracy__ERRANT__total_errors',
                          'accuracy__ERRANT__missing_errors','accuracy__ERRANT__unnecessary_errors']
                tables={}
                for method,order in [('mi',[0,1,2,3]),('ig',[0,2,1,3])]:
                    rows=[]
                    for target in ('cefr','level'):
                        for rank,index in enumerate(order,1):
                            row={'target':target,'feature':features[index],'eligible':True,
                                 ('mutual_information' if method=='mi' else 'information_gain_bits'):(5-rank)/10,
                                 'training_observations':100}
                            if method=='ig':row['best_split_threshold']=.25 if rank<4 else np.nan
                            rows.append({**row,'rank':rank,'selected':rank<=2})
                    tables[method]=pd.DataFrame(rows)
                settings=ranking_settings(SimpleNamespace(top_k=2,min_observations=2),manifest,int(splits.split.eq('train').sum()))
                source_hashes={method:hashlib.sha256(table.to_csv(index=False).encode()).hexdigest() for method,table in tables.items()}
                config={'format':'two_method_workbook','settings':settings,'source_table_sha256':source_hashes}
                with pd.ExcelWriter(path,engine='openpyxl') as writer:
                    for method,sheet in RANKING_SHEETS.items():tables[method].to_excel(writer,sheet_name=sheet,index=False)
                    metadata={'configuration':config,'configuration_id':token(config),
                              'table_sha256':{method:ranking_table_digest(table) for method,table in tables.items()}}
                    pd.DataFrame([[json.dumps(metadata)]]).to_excel(writer,sheet_name='_Pipeline_Metadata',index=False,header=False)
                    writer.sheets['_Pipeline_Metadata'].sheet_state='hidden'
                with patch.object(module,'load_matrix',side_effect=AssertionError('Do not recalculate saved rankings')):
                    migrated,actual=ranking_tables(folder,manifest)
                    main('cefr',['rank',*args,'--top-k','2'])
                self.assertEqual(migrated['source_table_sha256'],source_hashes)
                self.assertEqual(migrated['settings'],settings)
                for method in tables:pd.testing.assert_frame_equal(tables[method],actual[method])
                combined=pd.read_excel(path,sheet_name=RANKING_SHEET)
                self.assertEqual(len(combined),8)
                self.assertEqual(set(combined.selected_by),{'Both','MI only','IG only','Neither'})
                for target in ('cefr','level'):
                    definitions=experiment_definitions(SimpleNamespace(shared=root/'shared',ranking=folder,task=target,use_ranked=True),manifest)
                    self.assertEqual(definitions['training_mutual_information_features']['columns'],features[:2])
                    self.assertEqual(definitions['training_information_gain_features']['columns'],[features[0],features[2]])
                    self.assertEqual(definitions['training_mutual_information_features']['ranking_sha256'],source_hashes['mi'])
                    self.assertEqual(definitions['training_information_gain_features']['ranking_sha256'],source_hashes['ig'])
                book=openpyxl.load_workbook(path)
                self.assertEqual([sheet.title for sheet in book if sheet.sheet_state=='visible'],[RANKING_SHEET])
                self.assertEqual(book[RANKING_SHEET].freeze_panes,'C2')
                self.assertIn('TrainingFeatureRanking',book[RANKING_SHEET].tables)
                for row_number,values in enumerate(combined.sort_values(['target','ig_rank']).itertuples(index=False,name=None),2):
                    for column,value in enumerate(values,1):book[RANKING_SHEET].cell(row_number,column,value if pd.notna(value) else None)
                book.save(path);book.close()
                _,sorted_tables=ranking_tables(folder,manifest)
                for method in tables:pd.testing.assert_frame_equal(tables[method],sorted_tables[method])
                book=openpyxl.load_workbook(path)
                book[RANKING_SHEET]['L2']='Neither'
                book.save(path);book.close()
                with self.assertRaisesRegex(ValueError,'table changed'):ranking_tables(folder,manifest)
        def test_content_checks_survive_timestamp_changes(self):
            with tempfile.TemporaryDirectory() as temporary:
                root=Path(temporary); args=self.fixture(root)
                main('cefr',['prepare',*args])
                source=root/'data/feature_dictionary.csv'
                before=source.stat()
                os.utime(source, ns=(before.st_atime_ns, before.st_mtime_ns+1000000000))
                verify_shared(root/'shared')
                with source.open('a') as handle:handle.write('\n')
                with self.assertRaisesRegex(ValueError,'Input contents changed'):verify_shared(root/'shared')
        def test_legacy_setup_requires_explicit_rebuild(self):
            with tempfile.TemporaryDirectory() as temporary:
                root=Path(temporary);args=self.fixture(root)
                main('cefr',['prepare',*args])
                path=root/'shared/dataset_manifest.json';manifest=read_json(path)
                manifest['sources'][0].pop('sha256');write_json(path,manifest)
                with self.assertRaisesRegex(ValueError,'legacy timestamp'):verify_shared(root/'shared')
        def test_rebuild_preserves_previous_outputs(self):
            with tempfile.TemporaryDirectory() as temporary:
                root=Path(temporary);args=self.fixture(root)
                main('cefr',['prepare',*args])
                generated=root/'level/1_Intermediate_Calculations';generated.mkdir(parents=True)
                (generated/'previous_result.txt').write_text('retain this result')
                contributions=root/'level/3_Feature_Contributions/0_All_Experiments'
                contributions.mkdir(parents=True)
                (contributions/'feature_contributions.csv').write_text('retain explanations')
                wrapper=root/'level/level_prediction.py';wrapper.write_text('# retain code')
                with patch.object(module,'BASE',root):main('cefr',['prepare',*args,'--rebuild-setup'])
                backups=list((root/'9_Previous_Runs').glob('*/level/1_Intermediate_Calculations/previous_result.txt'))
                self.assertEqual(len(backups),1)
                self.assertEqual(backups[0].read_text(),'retain this result')
                self.assertEqual(len(list((root/'9_Previous_Runs').glob('*/level/3_Feature_Contributions/0_All_Experiments/feature_contributions.csv'))),1)
                self.assertTrue(wrapper.exists())
                verify_shared(root/'shared')
    result=unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(WorkflowTests))
    if not result.wasSuccessful():raise SystemExit(1)


if __name__=='__main__':
    main('cefr')
