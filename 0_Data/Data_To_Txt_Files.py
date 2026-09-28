import sys
from pathlib import Path
from openpyxl import load_workbook

script_dir = Path(__file__).parent
xlsx_path = script_dir / "Final database (main prompts).xlsx"

# Each job: (output_dir, source_text_column, label)
jobs = [
    (script_dir / "raw_texts", "text", "raw"),
    (script_dir / "corrected_texts", "text_corrected", "corrected"),
]

active_jobs = []
for output_dir, column, label in jobs:
    if output_dir.exists():
        print(f"'{output_dir.name}' already exists. Skipping {label} text creation.")
    else:
        output_dir.mkdir(exist_ok=True)
        active_jobs.append((output_dir, column, label))

if not active_jobs:
    print("Nothing to do — all output folders already exist.")
    sys.exit(0)

workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
worksheet = workbook.active

# Fix incorrect worksheet dimensions
worksheet.reset_dimensions()

rows = worksheet.iter_rows(values_only=True)

try:
    first_row = next(rows)
except StopIteration:
    raise ValueError("The worksheet is empty.")

headers = [
    str(value).strip().lower() if value is not None else ""
    for value in first_row
]

required_columns = {"writing_id", "cefr", "level", "grade"}
required_columns.update(column for _, column, _ in active_jobs)

missing_columns = required_columns - set(headers)

if missing_columns:
    raise ValueError(
        f"Missing columns: {sorted(missing_columns)}\n"
        f"Available columns: {headers}"
    )

files_created = {label: 0 for _, _, label in active_jobs}

for values in rows:
    row = dict(zip(headers, values))

    writing_id = str(row.get("writing_id") or "").strip()
    cefr = str(row.get("cefr") or "").strip()
    level = str(row.get("level") or "").strip()
    grade = str(row.get("grade") or "").strip()

    if not writing_id:
        continue

    filename = (
        f"{writing_id}_cefr_{cefr}_level_{level}_grade_{grade}.txt"
    )

    for output_dir, column, label in active_jobs:
        text = str(row.get(column) or "").strip()
        with (output_dir / filename).open("w", encoding="utf-8") as file:
            file.write(text)
        files_created[label] += 1

workbook.close()

for output_dir, _, label in active_jobs:
    print(f"\nCreated {files_created[label]} {label} text files.")
    print(f"Output folder: {output_dir}")